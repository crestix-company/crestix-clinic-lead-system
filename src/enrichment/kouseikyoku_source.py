"""関東信越厚生局の帳票Excelを1施設1行に変換するアダプター。"""
from io import BytesIO
from pathlib import Path
from urllib.parse import urljoin
import re
import zipfile
import pandas as pd
from openpyxl import load_workbook
from bs4 import BeautifulSoup
import requests
from src.io.input_loader import load_table
from src.utils.date_utils import parse_date

OFFICIAL_INDEX = "https://kouseikyoku.mhlw.go.jp/kantoshinetsu/chousa/shitei.html"
MASTER_ALIASES = {
    "clinic_id": ["clinic_id", "医療機関番号", "医療機関コード", "ID"],
    "clinic_name": ["clinic_name", "医療機関名", "医療機関名称", "医院名", "施設名"],
    "address": ["address", "所在地", "医療機関所在地", "住所"],
    "phone": ["phone", "電話番号", "TEL", "電話"],
    "prefecture": ["prefecture", "都道府県"],
    "medical_type": ["medical_type", "医科/歯科区分", "医科歯科区分"],
    "facility_type": ["facility_type", "病院/診療所区分", "病院診療所区分"],
    "owner_name": ["owner_name", "開設者氏名", "開設者名"],
    "manager_name": ["manager_name", "管理者氏名", "院長名", "管理者名"],
    "designation_date": ["designation_date", "指定年月日"],
    "manager_age": ["manager_age", "管理者年齢", "院長年齢"],
    "manager_birth_year": ["manager_birth_year", "管理者生年", "院長生年"],
    "license_registration_year": ["license_registration_year", "医籍登録年", "医師免許取得年"],
    "graduation_year": ["graduation_year", "大学卒業年", "医学部卒業年"],
    "registration_reason": ["registration_reason", "登録理由"],
    "status": ["status", "営業状況", "現在営業中", "現存休止", "廃止/休止の有無"],
    "as_of": ["as_of", "データ基準日"],
    "source_url": ["source_url", "出典URL"],
    "departments": ["departments", "診療科", "診療科目", "標榜科目"],
}


def canonical_master(table):
    result = pd.DataFrame(index=table.data.index)
    for field, aliases in MASTER_ALIASES.items():
        hits = [i for i, name in enumerate(table.headers) if name.strip() in aliases]
        if len(hits) > 1:
            raise ValueError(f"マスタの{field}候補列が複数あります。見出しを整理してください。")
        result[field] = table.data[hits[0]] if hits else ""
    for i, name in enumerate(table.headers):
        if name not in result.columns and name not in sum(MASTER_ALIASES.values(), []):
            result[name] = table.data[i]
    if not result["clinic_name"].astype(bool).any() and not result["phone"].astype(bool).any():
        raise ValueError("マスタに医院名または電話番号列が必要です。サンプルの見出しを使ってください。")
    return result.fillna("").astype(str)


def parse_kanto_excel(raw, prefecture="東京都", source_url="", sheet=None):
    book = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    ws = book[sheet] if sheet else book.active
    rows = ws.iter_rows(values_only=True)
    records, record, basis, title_seen = [], None, "", False
    medical_type = "医科"
    def department_text(value):
        # 無床診療所は施設の先頭行、有床施設は続きの行に診療科がある。
        text = re.sub(r"(?:一般|療養|精神|感染|結核)\s*[0-9０-９]+", "", value).strip()
        return "" if re.fullmatch(r"[\s0-9０-９]*", text) else text
    def commit():
        if record:
            records.append(record.copy())
    for raw_row in rows:
        cells = [str(v or "").strip() for v in raw_row]
        cells += [""] * max(0, 10-len(cells))
        if "コード内容別医療機関一覧表" in cells[0]:
            title_seen = True
        if not basis and "現在" in cells[0]:
            medical_type = "歯科" if "歯科" in cells[0] else "医科"
            match = re.search(r"((?:令和|平成)\s*\d+年\s*\d+月\s*\d+日)現在", cells[0])
            if match:
                d = parse_date(match[1])
                basis = d.isoformat() if d else ""
        if cells[0].isdigit() and re.fullmatch(r"\d{2},\d{4},\d", cells[1]):
            commit()
            designation = parse_date(cells[7])
            address = re.sub(r"^〒[0-9０-９－\-]+\s*", "", cells[3])
            if not address.startswith(prefecture):
                address = prefecture + address
            record = {
                "clinic_id": prefecture + ":" + cells[1].replace(",", ""),
                "clinic_name": cells[2], "address": address, "phone": cells[4],
                "prefecture": prefecture, "medical_type": medical_type, "facility_type": "病院" if cells[9] == "特定機能" else cells[9],
                "owner_name": cells[5], "manager_name": cells[6],
                "designation_date": designation.isoformat() if designation else cells[7],
                "registration_reason": "", "status": "", "as_of": basis,
                "source_url": source_url, "designation_period_start": "", "departments": department_text(cells[8]),
            }
        elif record and not cells[0] and not cells[1]:
            if department_text(cells[8]):
                record["departments"] = (record["departments"] + " " + department_text(cells[8])).strip()
            if cells[9] in {"現存", "休止", "廃止", "辞退", "取消"}:
                record["status"] = cells[9]
            if cells[7]:
                parsed = parse_date(cells[7])
                if parsed:
                    record["designation_period_start"] = parsed.isoformat()
                else:
                    record["registration_reason"] = cells[7]
    commit()
    book.close()
    if not title_seen or not records or not basis:
        raise ValueError("関東信越厚生局Excelの形式を確認できません。標準CSVマスタに変換して読み込んでください。")
    if any(not r["clinic_name"] or r["facility_type"] not in {"病院", "診療所"} for r in records):
        raise ValueError("帳票の列位置または区分が想定と異なります。自動取込を停止しました。")
    return pd.DataFrame(records).fillna("")


def load_master(raw, filename, mode="standard", prefecture="東京都", sheet=None):
    if mode == "kanto_excel":
        return parse_kanto_excel(raw, prefecture, OFFICIAL_INDEX, sheet)
    return canonical_master(load_table(raw, filename, sheet=sheet))


def download_tokyo_master(timeout=30):
    """利用者が明示操作したときだけ実行。403/429等は再試行せず失敗を返す。"""
    session = requests.Session()
    session.headers["User-Agent"] = "ClinicListFilter/0.1 (public-data import)"
    response = session.get(OFFICIAL_INDEX, timeout=timeout)
    response.raise_for_status()
    soup = BeautifulSoup(response.content, "html.parser")
    hrefs = [a.get("href", "") for a in soup.select("a[href]")]
    href = next((h for h in hrefs if re.search(r"/shitei_ika_[^/]+\.zip$", h)), None)
    if not href:
        raise ValueError("医科Excel ZIPのリンクを確認できません。公式ページから手動でダウンロードしてください。")
    url = urljoin(OFFICIAL_INDEX, href)
    if not url.startswith("https://kouseikyoku.mhlw.go.jp/kantoshinetsu/"):
        raise ValueError("公式配布先を確認できません。")
    download = session.get(url, timeout=timeout)
    download.raise_for_status()
    with zipfile.ZipFile(BytesIO(download.content)) as archive:
        entries = [i for i in archive.infolist() if "東京" in i.filename and i.filename.endswith(".xlsx")]
        if len(entries) != 1 or entries[0].file_size > 50*1024**2:
            raise ValueError("東京都Excelが一意でないか、大きすぎます。手動取込をご利用ください。")
        raw = archive.read(entries[0])
        master = parse_kanto_excel(raw, "東京都", url)
        return master, raw, Path(entries[0].filename).name
