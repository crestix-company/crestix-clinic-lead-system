from dataclasses import dataclass, field
from datetime import datetime, date
from io import BytesIO, StringIO
from pathlib import Path
import csv
import re
import unicodedata
import pandas as pd
from openpyxl import load_workbook
from src.normalizer.phone import normalize_phone

ALIASES = {
    "phone": ["電話番号", "医院電話番号", "TEL", "電話", "phone", "telephone"],
    "clinic_name": ["医院名", "医療機関名", "医療機関名称", "クリニック名", "施設名", "会社名", "clinic_name"],
    "address": ["住所", "所在地", "医療機関所在地", "address"],
    "url": ["URL", "HP", "ホームページ", "公式サイト", "website"],
    "epark_url": ["EPARK URL", "EPARK_URL", "EPARK", "epark_url"],
}


@dataclass
class InputTable:
    headers: list[str]
    data: pd.DataFrame  # 列位置をキーにする。重複・空の見出しも壊さない。
    encoding: str = "utf-8-sig"
    delimiter: str = ","
    sheet_name: str = ""
    warnings: list[str] = field(default_factory=list)

    def value(self, row, column):
        return self.data.iat[row, column] if column is not None else ""

    def frame(self):
        result = self.data.copy()
        result.columns = self.headers
        return result


def decode_csv(raw, encoding="auto"):
    if b"\x00" in raw:
        raise ValueError("UTF-16など未対応の文字コードです。UTF-8またはCP932で保存してください。")
    encodings = ["utf-8-sig", "cp932", "shift_jis"] if encoding == "auto" else [encoding]
    for enc in encodings:
        try:
            return raw.decode(enc, errors="strict"), enc
        except UnicodeDecodeError:
            continue
    raise ValueError("文字コードを判別できません。文字コードを選択するか、UTF-8で保存し直してください。")


def sheet_names(raw, filename):
    if Path(filename).suffix.lower() not in {".xlsx", ".xls"}:
        return []
    with pd.ExcelFile(BytesIO(raw)) as book:
        return book.sheet_names


def _string(value, number_format=""):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (float, int)):
        if float(value).is_integer():
            integer = str(int(value))
            if re.fullmatch(r"0+", number_format or ""):
                return integer.zfill(len(number_format))
            return integer
    return str(value)


def load_table(raw, filename, encoding="auto", sheet=None, header_row=0, max_rows=100000, max_bytes=50*1024**2):
    if len(raw) > max_bytes:
        raise ValueError("ファイルサイズが上限を超えています。")
    suffix = Path(filename).suffix.lower()
    warnings = []
    enc, delimiter, selected = "utf-8-sig", ",", ""
    if suffix == ".csv":
        text, enc = decode_csv(raw, encoding)
        try:
            delimiter = csv.Sniffer().sniff(text[:20000], delimiters=",;\t").delimiter
        except csv.Error:
            delimiter = ","
        try:
            rows = list(csv.reader(StringIO(text, newline=""), delimiter=delimiter, strict=True))
        except csv.Error as exc:
            raise ValueError("CSVの引用符や改行が不正です。元ファイルをご確認ください。") from exc
        rows = rows[header_row:]
    elif suffix == ".xlsx":
        try:
            book = load_workbook(BytesIO(raw), read_only=True, data_only=False)
            selected = sheet if sheet is not None else book.sheetnames[0]
            if selected not in book.sheetnames:
                raise ValueError("選択したシートがありません。")
            ws = book[selected]
            if ws.max_row > max_rows + header_row + 1:
                raise ValueError("行数が上限を超えています。")
            rows = []
            for cells in ws.iter_rows(min_row=header_row + 1):
                if any(c.data_type == "f" for c in cells):
                    raise ValueError("数式セルがあります。Excelで対象シートを値貼り付けしてからアップロードしてください。")
                rows.append([_string(c.value, c.number_format) for c in cells])
            book.close()
            # 末尾の空行だけを除去。空見出し・書式付きの末尾列も列構成として保持する。
            while rows and not any(rows[-1]):
                rows.pop()
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Excelを読み込めません。暗号化・破損・拡張子を確認してください。") from exc
        warnings.append("Excelの値を読み込みました。書式・数式・他シートは最終CSVには含みません。")
    elif suffix == ".xls":
        try:
            import xlrd
            book = xlrd.open_workbook(file_contents=raw, formatting_info=True)
            selected = sheet if sheet is not None else book.sheet_names()[0]
            ws = book.sheet_by_name(selected)
            rows = []
            for r in range(header_row, ws.nrows):
                values = []
                for c in range(ws.ncols):
                    cell = ws.cell(r, c)
                    fmt = book.format_map[book.xf_list[cell.xf_index].format_key].format_str
                    value = xlrd.xldate.xldate_as_datetime(cell.value, book.datemode) if cell.ctype == xlrd.XL_CELL_DATE else cell.value
                    values.append(_string(value, fmt))
                rows.append(values)
            warnings.append("旧形式Excelは保存済みの計算結果を読み込みます。Excelで再計算・保存済みかご確認ください。")
        except Exception as exc:
            raise ValueError("旧形式Excelを読み込めません。xlsxまたはCSVに保存し直してください。") from exc
    else:
        raise ValueError("CSV・XLSX・XLSを指定してください。")
    if not rows or not rows[0]:
        raise ValueError("ヘッダー行が見つかりません。")
    headers, values = rows[0], rows[1:]
    if len(values) > max_rows:
        raise ValueError("行数が上限を超えています。分割してアップロードしてください。")
    for i, row in enumerate(values, 2 + header_row):
        if len(row) != len(headers):
            raise ValueError(f"{i}行目の列数が見出しと一致しません。空行またはCSVの区切りを確認してください。")
    if len(set(headers)) != len(headers):
        warnings.append("重複した列名があります。列番号で区別し、出力にもそのまま保持します。")
    return InputTable(headers, pd.DataFrame(values, columns=range(len(headers)), dtype=str), enc, delimiter, selected, warnings)


def infer_columns(table):
    def key(s):
        return re.sub(r"[\s_]", "", unicodedata.normalize("NFKC", s)).casefold()
    result = {}
    for field, names in ALIASES.items():
        found = [i for i, h in enumerate(table.headers) if key(h) in {key(x) for x in names}]
        result[field] = found[0] if len(found) == 1 else None
    if result["phone"] is None and len(table.data):
        candidates = []
        for col in table.data.columns:
            vals = [str(v) for v in table.data[col].head(100) if str(v)]
            if vals and sum(bool(re.fullmatch(r"0\d{9,10}", normalize_phone(v))) for v in vals)/len(vals) >= .9:
                candidates.append(col)
        if len(candidates) == 1:
            result["phone"] = candidates[0]
    return result
