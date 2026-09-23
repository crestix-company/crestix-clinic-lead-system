from io import BytesIO, StringIO
from pathlib import Path
import csv
import math
import json
import unicodedata
import zipfile
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule


def csv_bytes(headers, rows, encoding="utf-8-sig", delimiter=","):
    stream = StringIO(newline="")
    writer = csv.writer(stream, delimiter=delimiter, lineterminator="\r\n")
    writer.writerow(headers)
    for row in rows:
        writer.writerow(["" if v is None or (isinstance(v, float) and math.isnan(v)) else v for v in row])
    try:
        return stream.getvalue().encode(encoding, errors="strict")
    except UnicodeEncodeError as exc:
        raise ValueError("選択した文字コードで出力できない文字があります。UTF-8を選択してください。") from exc


def final_csv(table, indices, encoding=None):
    # 正規化済み値や判定列を使わず、元の行・元の列位置だけで出力。
    return csv_bytes(table.headers, table.data.iloc[indices].itertuples(index=False, name=None),
                     encoding or table.encoding, table.delimiter)


def judged_frame(table, judgments):
    extra = judgments.copy()
    used = set(table.headers)
    names = []
    for original in extra.columns:
        name = original
        while name in used:
            name = "[判定]" + name
        names.append(name)
        used.add(name)
    extra.columns = names
    return pd.concat([table.frame().reset_index(drop=True), extra.reset_index(drop=True)], axis=1)


def _excel_value(value):
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    if hasattr(value, "item"):
        return value.item()
    return value


def xlsx_bytes(sheets, probabilities=()):
    book = Workbook()
    book.remove(book.active)
    for name, frame in sheets.items():
        ws = book.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "C2" if len(frame.columns) > 3 else "A2"
        for r, values in enumerate([list(frame.columns)] + list(frame.itertuples(index=False, name=None)), 1):
            for c, value in enumerate(values, 1):
                value = _excel_value(value)
                cell = ws.cell(r, c, value)
                # ユーザー入力が = で始まっても数式として実行しない。値自体は不変。
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.font = Font(name="Yu Gothic", size=10, color="172B4D")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if r == 1:
                    cell.fill = PatternFill("solid", fgColor="243B53")
                    cell.font = Font(name="Yu Gothic", size=10, bold=True, color="FFFFFF")
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                elif str(frame.columns[c-1]) in probabilities:
                    cell.number_format = "0.0%"
            if r == 1:
                ws.row_dimensions[r].height = 32
        for c, heading in enumerate(frame.columns, 1):
            width = 20
            if any(k in str(heading) for k in ["医院名", "住所", "理由", "出典", "URL", "特徴", "設定"]):
                width = 48
            if name == "判定条件":
                width = 80 if c == 1 else 100
            ws.column_dimensions[get_column_letter(c)].width = width
        for r in range(2, len(frame)+2):
            lines = 1
            for c in range(1, len(frame.columns)+1):
                text = str(ws.cell(r, c).value or "")
                capacity = max(8, ws.column_dimensions[get_column_letter(c)].width-2)
                count = sum(max(1, math.ceil(sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in line)/capacity))
                            for line in text.split("\n"))
                lines = max(lines, count)
            ws.row_dimensions[r].height = min(409, 16*lines + 6)
        if name != "判定条件":
            ws.auto_filter.ref = ws.dimensions
        if "営業優先度" in frame.columns:
            col = get_column_letter(list(frame.columns).index("営業優先度")+1)
            for label, color in [("S", "DCFCE7"), ("A", "DBEAFE"), ("B", "E0E7FF"), ("C", "FEF3C7"), ("除外", "FEE2E2")]:
                ws.conditional_formatting.add(f"{col}2:{col}{max(2,len(frame)+1)}",
                    CellIsRule(operator="equal", formula=[f'"{label}"'], fill=PatternFill("solid", fgColor=color)))
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def build_outputs(table, result, encoding=None):
    judged = judged_frame(table, result.judgments)
    output_encoding = encoding or table.encoding
    probabilities = [c for c in judged.columns if "59歳以下確率" in c or "医籍信頼度" in c]
    condition_rows = []
    def flatten(prefix, value):
        if isinstance(value, dict):
            for key, child in value.items():
                flatten(prefix+" / "+str(key), child)
        elif isinstance(value, list):
            condition_rows.append((prefix, json.dumps(value, ensure_ascii=False)))
        else:
            condition_rows.append((prefix, value))
    for name, value in result.metadata.items():
        if name in {"条件設定", "年齢モデル", "判定設定"}:
            flatten(name, json.loads(value))
        else:
            flatten(name, value)
    conditions = pd.DataFrame(condition_rows + list(result.metrics.items()), columns=["項目", "内容"])
    files = {
        "final_comdesk_import.csv": final_csv(table, result.final_indices, output_encoding),
        "judged_all.xlsx": xlsx_bytes({"判定付き全件": judged, "判定条件": conditions}, probabilities),
        "excluded.csv": csv_bytes(judged.columns, judged.iloc[result.excluded_indices].itertuples(index=False, name=None), output_encoding),
        "needs_review.csv": csv_bytes(judged.columns, judged.iloc[result.review_indices].itertuples(index=False, name=None), output_encoding),
    }
    if table.sheet_name:
        files["final_comdesk_import.xlsx"] = xlsx_bytes({table.sheet_name[:31]: table.frame().iloc[result.final_indices]})
    return files


def zip_outputs(files):
    stream = BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return stream.getvalue()


def write_outputs(files, directory):
    path = Path(directory)
    path.mkdir(parents=True, exist_ok=True)
    for name, data in files.items():
        (path / name).write_bytes(data)
