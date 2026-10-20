"""Merge connector-read-only Phase I Sheets ranges into the local Phase G workbook."""
from __future__ import annotations

import json
import os
from pathlib import Path
import sys

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.outreach.google_sheets_readonly_adapter import GoogleSheetsReadOnlyAdapter


BASE = Path("/private/tmp")
SOURCE_JSON = Path(os.environ.get("PHASE_I_SHEETS_JSON_PATH", str(BASE / "clinic_phase_i_sheets_readonly.json")))
BASE_WORKBOOK = Path(os.environ.get("PHASE_G_SHEETS_XLSX_PATH", str(BASE / "clinic_phase_g_source.xlsx")))
OUTPUT_WORKBOOK = Path(os.environ.get("PHASE_I_SHEETS_XLSX_PATH", str(BASE / "clinic_phase_i_source.xlsx")))
MANIFEST = BASE / "clinic_phase_i_sheets_manifest.json"
TAB_NAMES = (
    "案件管理", "メール管理", "トラッキングログ", "日程候補ログ",
    "328医院_現在進捗_20261018",
)


def build_snapshot() -> dict[str, object]:
    payload = json.loads(SOURCE_JSON.read_text(encoding="utf-8"))
    raw = payload["tabs"]
    by_range = {f"'{name}'!A:ZZ": [tab["headers"], *tab["rows"]]
                for name, tab in raw.items()}
    reader = GoogleSheetsReadOnlyAdapter(lambda spreadsheet_id, a1: by_range.get(a1, ()))
    read_tabs = reader.read_tabs(str(payload["spreadsheet_id"]), TAB_NAMES)

    workbook = load_workbook(BASE_WORKBOOK)
    counts: dict[str, int] = {}
    for tab in read_tabs:
        if tab.tab_name not in workbook.sheetnames:
            raise ValueError(f"Source tab is absent from the base snapshot: {tab.tab_name}")
        sheet = workbook[tab.tab_name]
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cell.value = None
        matrix = [list(tab.headers), *[list(row.values()) for row in tab.rows]]
        for row_index, values in enumerate(matrix, start=1):
            for col_index, value in enumerate(values, start=1):
                sheet.cell(row_index, col_index).value = value
        counts[tab.tab_name] = len(tab.rows)

    workbook.save(OUTPUT_WORKBOOK)
    os.chmod(OUTPUT_WORKBOOK, 0o600)
    manifest = {
        "snapshot_type": payload["snapshot_type"],
        "spreadsheet_title": payload["spreadsheet_title"],
        "retrieved_at": payload["retrieved_at"],
        "tab_counts": counts,
        "source_adapter": "GoogleSheetsReadOnlyAdapter.read_tabs",
        "writes_performed": False,
    }
    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    os.chmod(MANIFEST, 0o600)
    return manifest


if __name__ == "__main__":
    print(json.dumps(build_snapshot(), ensure_ascii=False))
