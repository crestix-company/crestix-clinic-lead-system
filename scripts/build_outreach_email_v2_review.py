"""Build a protected, local v2 email review workbook from the immutable prior cohort.

This script is preview-only. It does not access DB/Gmail/network and never sends or drafts mail.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import os
import stat
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from src.outreach.email_template_v2 import (
    ARTICLE_ABSENT, ARTICLE_PRESENT, ARTICLE_UNVERIFIED, PreviewInput,
    load_settings, render_preview,
)

EXPECTED_SHA256 = "9608e4ac5ecac98af18c29f7375a71d1232b3b3c90c7b9c337d5bbe9df4b07fa"
DEFAULT_SOURCE = Path("artifacts/outreach_phase12_20261016_v2_evidence_review/outreach_email_review_takahara_20261016_v2_evidence_review.xlsx")
DEFAULT_MASTER = Path(".codex-input/article_selection_v02/原本_変更禁止_Crestix_クリニックナビ記事一覧.csv")
DEFAULT_OUTPUT = Path("artifacts/outreach_phase12_20261016_v2_final_r2")
EXPECTED_COUNT = 329


def approved_v2_body(preview, *, human_review_status: str) -> str:
    """Return the sole delivery-source body; never fall back to the legacy Phase 2 body."""
    if human_review_status.strip() != "承認" or not preview.ready_for_draft:
        return ""
    return preview.body


def _csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def build(source: Path, master: Path, output_dir: Path, settings_path: Path) -> Path:
    if hashlib.sha256(master.read_bytes()).hexdigest() != EXPECTED_SHA256:
        raise RuntimeError("Immutable article master SHA-256 mismatch")
    master_rows = _csv_rows(master)
    article_urls = frozenset((r.get("URL") or "").strip() for r in master_rows if (r.get("URL") or "").strip())
    if len(master_rows) != 4000:
        raise RuntimeError("Unexpected immutable article master row count")
    settings = load_settings(settings_path)

    wb = load_workbook(source)
    ws = wb["全件一覧"]
    headers = [c.value for c in ws[1]]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    if len(rows) != EXPECTED_COUNT:
        raise RuntimeError(f"Expected preserved cohort of {EXPECTED_COUNT}; found {len(rows)}")
    idx = {name: i for i, name in enumerate(headers)}
    required = {"clinic_id", "Clinic UUID", "Email Candidate ID", "医院名", "院長名（Clinic Master記録）",
                "メールアドレス", "記事判定", "実掲載記事URL", "マイナビ医院詳細ページURL",
                "AI第一推奨URL", "人間確認ステータス", "修正コメント", "最終確認日"}
    if not required.issubset(idx):
        raise RuntimeError("Input workbook does not have required preserved columns")

    new_headers = [
        "v2件名プレビュー", "v2本文プレビュー", "本文プレビュー状態", "Ready For Draft",
        "記事判定証拠（詳細ページ）", "記事判定証拠（紹介記事）", "記事確認日時・方法",
        "院長名の利用可否", "候補日時設定", "候補日時の利用可否", "電話履歴確認",
        "v2生成ブロッカー",
    ]
    start_col = ws.max_column + 1
    for c, header in enumerate(new_headers, start_col):
        ws.cell(1, c, header)
        ws.column_dimensions[__import__("openpyxl").utils.get_column_letter(c)].width = 72 if "理由" in header else 42
    run_date = __import__("datetime").date.today()
    count_status: dict[str, int] = {}
    body_count = ready_count = 0
    fujita_rows = haihara_rows = 0
    for rno, values in enumerate(rows, start=2):
        row = dict(zip(headers, values))
        clinic_name = str(row.get("医院名") or "").strip()
        article_status = str(row.get("記事判定") or ARTICLE_UNVERIFIED)
        if article_status not in {ARTICLE_PRESENT, ARTICLE_ABSENT, ARTICLE_UNVERIFIED}:
            article_status = ARTICLE_UNVERIFIED
        actual_url = str(row.get("実掲載記事URL") or "").strip()
        profile_url = str(row.get("マイナビ医院詳細ページURL") or "").strip()
        recommendation_url = str(row.get("AI第一推奨URL") or "").strip()
        # Only the independently recorded Fuji article/profile evidence is considered verified.
        evidence_detail = str(row.get("記事判定根拠URL（詳細ページ）") or "").strip()
        evidence_article = str(row.get("記事判定根拠URL（紹介記事）") or "").strip()
        evidence_method = str(row.get("確認方法・根拠") or "").strip()
        if article_status == ARTICLE_PRESENT:
            verified_listing = bool(evidence_detail and evidence_article == actual_url and actual_url in article_urls)
        else:
            verified_listing = False
        director_raw = str(row.get("院長名（Clinic Master記録）") or "").strip()
        review_status = str(row.get("人間確認ステータス") or "未確認").strip()
        # The source column only identifies a Clinic Master manager_name value, not director proof.
        settings_dates = tuple(settings.get("candidate_dates", ()))
        preview = render_preview(PreviewInput(
            clinic_name=clinic_name,
            director_name="",
            director_name_verified=False,
            article_status=article_status,
            actual_article_url=actual_url,
            recommended_article_url=recommendation_url,
            clinic_detail_url=profile_url,
            clinic_detail_verified=bool(evidence_detail),
            article_listing_verified=verified_listing,
            article_absence_verified=False,
            schedule=settings_dates,
            sender_name=settings["sender_name"],
            phone_call_date=None,
            email_send_request_confirmed=False,
            recipient_email_verified=bool(str(row.get("メールアドレス") or "").strip()),
            gmail_account_verified=False,
            human_approved=review_status == "承認",
            duplicate_check_complete=False,
            article_master_urls=article_urls,
            as_of=run_date,
        ))
        if clinic_name == "ふじた矯正歯科クリニック":
            fujita_rows += 1
        if clinic_name == "はいはら歯科":
            haihara_rows += 1
        reason_parts = list(preview.blockers)
        if article_status == ARTICLE_UNVERIFIED:
            reason_parts.append("医院詳細ページ・紹介記事の対応証拠未確認")
        if article_status == ARTICLE_ABSENT:
            reason_parts.append("紹介記事がない証拠未確認。記事なし用本文を生成しない")
        if not director_raw:
            reason_parts.append("Clinic Master記録にも院長名なし")
        else:
            reason_parts.append("Clinic Master記録名はあるが院長本人との一致・出典未確認のため差し込み禁止")
        reason_parts.append("候補日時は設定ファイル未設定" if not settings_dates else "候補日時は本人予定と未照合")
        if not verified_listing:
            reason_parts.append("記事掲載証拠なし、または原本URL完全一致未確認")
        status = "PREVIEW_GENERATED_NEEDS_REVIEW" if preview.preview_generated else "BODY_NOT_GENERATED_NEEDS_REVIEW"
        # Column X is historical only and is intentionally never read as a fallback.
        final_body = approved_v2_body(
            preview,
            human_review_status=review_status,
        )
        ws.cell(rno, 22, final_body)  # V: only approved, fully gated v2 body
        ws.cell(rno, start_col, preview.subject if preview.preview_generated else "")
        ws.cell(rno, start_col + 1, preview.body if preview.preview_generated else "")
        ws.cell(rno, start_col + 2, status)
        ws.cell(rno, start_col + 3, "NO")
        ws.cell(rno, start_col + 4, evidence_detail)
        ws.cell(rno, start_col + 5, evidence_article)
        ws.cell(rno, start_col + 6, evidence_method)
        ws.cell(rno, start_col + 7, "未確認（Clinic Masterの管理者名は差し込みに不使用）")
        ws.cell(rno, start_col + 8, "、".join(settings_dates))
        ws.cell(rno, start_col + 9, "NO" if not settings_dates else "NO（カレンダー照合なし）")
        ws.cell(rno, start_col + 10, "未確認（中立導入のみ）")
        ws.cell(rno, start_col + 11, "／".join(dict.fromkeys(x for x in reason_parts if x)))
        count_status[status] = count_status.get(status, 0) + 1
        body_count += int(preview.preview_generated)
        ready_count += int(preview.ready_for_draft)
        for c in (start_col, start_col + 1, start_col + 11):
            ws.cell(rno, c).alignment = Alignment(wrap_text=True, vertical="top")
        for c in (start_col + 2, start_col + 3, start_col + 7, start_col + 9, start_col + 10):
            ws.cell(rno, c).fill = PatternFill("solid", fgColor="FCE4D6")
    if fujita_rows != 1:
        raise RuntimeError("The recorded Fuji evidence row is missing or duplicated")

    # Maintain original review statuses/comments; add explicit formula-free summary and examples.
    summary = wb.create_sheet("v2集計・テンプレート例")
    summary.append(["項目", "値・理由"])
    summary_rows = [
        ("対象母集団", EXPECTED_COUNT),
        ("本文プレビュー生成", body_count),
        ("本文未生成", EXPECTED_COUNT - body_count),
        ("記事判定 ARTICLE_PRESENT", sum(str(r[idx["記事判定"]]) == ARTICLE_PRESENT for r in rows)),
        ("記事判定 ARTICLE_ABSENT", sum(str(r[idx["記事判定"]]) == ARTICLE_ABSENT for r in rows)),
        ("記事判定 ARTICLE_UNVERIFIED", sum(str(r[idx["記事判定"]]) == ARTICLE_UNVERIFIED for r in rows)),
        ("Ready For Draft", ready_count),
        ("正式テンプレート送信者", "高原（固定。髙原へ自動変換しない）"),
        ("候補日時", "設定なし。2026-10-19〜23を既定値にしていません。" if not settings_dates else "設定値は本人予定と未照合。"),
        ("確認状態保持", "入力Excelの人間確認ステータス・修正コメント・最終確認日は変更なし。"),
        ("テンプレートv2", "正式本文はsrc/outreach/email_template_v2.pyに実装。記事あり/なし、架電未確認の安全分岐を実装。"),
        ("A: ふじた矯正歯科クリニック", "ARTICLE_PRESENTの公開根拠を継承。本文プレビューは院長名・候補日時プレースホルダー付き、送信不可。"),
        ("ARTICLE_ABSENTテンプレート例（実医院判定ではない）", """{医院名}
院長 {院長名} 様

お世話になっております。
株式会社Crestixの高原と申します。

本日、貴院へお電話させていただき、ご対応いただいた方よりメールにてお送りするよう伺いましたので、ご連絡いたしました。

現在、マイナビクリニックナビにて掲載されている貴院の医院情報について、掲載データの共有・現状のご報告と、掲載内容の確認をさせていただければと考えております。

【記事例】
{AI第一推奨記事URL}

【掲載いただいているページ】
{医院詳細ページURL}

なお、掲載内容についての行き違い防止のため、お打ち合わせ当日は院長先生にもご参加いただく形でお願いしております。

恐れ入りますが、院長先生がご参加可能な日時を、下記候補よりお知らせいただけますでしょうか。

{候補日時}

上記よりオンラインにて20分～30分ほど頂戴できますと幸いです。
何卒よろしくお願い申し上げます。"""),
        ("B: はいはら歯科", "329医院の対象外。公式検索ではMynavi地域一覧で医院名を確認できるが、医院詳細URLと紹介記事タブの不掲載証拠は確認できないためARTICLE_ABSENTにせず本文生成しない。"),
        ("OAuth/Gmail", "認可・Gmail下書き・送信は未実施。OAuth準備は別途実装状態を監査。"),
    ]
    for item in summary_rows:
        summary.append(list(item))
    summary.freeze_panes = "A2"
    summary.auto_filter.ref = f"A1:B{summary.max_row}"
    summary.column_dimensions["A"].width = 40
    summary.column_dimensions["B"].width = 120
    for cell in summary[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(bold=True, color="FFFFFF")
    for rno in range(2, summary.max_row + 1):
        summary.cell(rno, 2).alignment = Alignment(wrap_text=True, vertical="top")

    if output_dir.exists():
        raise FileExistsError(f"Refusing to overwrite existing output directory: {output_dir}")
    output_dir.mkdir(parents=True, mode=0o700)
    output_dir.chmod(0o700)
    output = output_dir / "outreach_email_review_takahara_20261016_v2_final.xlsx"
    if output.exists():
        raise FileExistsError("Refusing to overwrite existing Excel output")
    wb.save(output)
    output.chmod(0o600)
    wb.close()
    if stat.S_IMODE(output_dir.stat().st_mode) != 0o700 or stat.S_IMODE(output.stat().st_mode) != 0o600:
        raise RuntimeError("Output file/directory permissions are not private")
    return output


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--master", type=Path, default=DEFAULT_MASTER)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--settings", type=Path, default=Path("config/outreach_email_v2.json"))
    args = parser.parse_args()
    output = build(args.source, args.master, args.output_dir, args.settings)
    check = load_workbook(output, read_only=True, data_only=True)
    ws = check["全件一覧"]
    print(f"rows={ws.max_row - 1} preview_rows={sum(1 for r in ws.iter_rows(min_row=2, values_only=True) if r[-10] == 'PREVIEW_GENERATED_NEEDS_REVIEW')} draft_ready=0")
    print(f"sheets={','.join(check.sheetnames)} output={output}")
    check.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
