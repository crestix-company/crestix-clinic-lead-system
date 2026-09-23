from io import BytesIO
from pathlib import Path
import json
import zipfile

import pandas as pd
from openpyxl import Workbook

from src.master.google_maps import QUEUE_HEADERS
from src.master.national_maps_queue import (
    MASTER_HEADERS, OFFICIAL_SOURCES, PREFECTURES, OfficialSource, build_maps_population, build_package,
    parse_code_workbook, split_three,
)


def make_book(pref="大阪府"):
    book = Workbook(); ws = book.active
    ws.title = pref
    ws.append(["コード内容別医療機関一覧表"])
    ws.append(["[令和 8年 9月 1日現在　医科　現存/休止]"])
    ws.append(["1", "01,1234,5", "見本クリニック", "〒530-0001 大阪市北区梅田1-1", "06-1234-5678\n常勤:1", "開設 太郎", "院長 花子", "令2.1.1\n新規\n令2.1.1", "内 消", "診療所\n現存"])
    ws.append(["2", "01,1234,6", "見本病院", "〒530-0002 大阪市北区梅田1-2", "06-1234-5679", "A", "B", "令2.1.1", "一般 20 内", "病院\n現存"])
    ws.append(["3", "01,1234,7", "見本医療センター", "〒530-0003 大阪市北区梅田1-3", "06-1234-5680", "A", "B", "令2.1.1", "内", "診療所\n現存"])
    ws.append(["4", "01,1234,8", "休止クリニック", "〒530-0004 大阪市北区梅田1-4", "06-1234-5681", "A", "B", "令2.1.1", "内", "診療所\n休止"])
    out = BytesIO(); book.save(out); return out.getvalue()


def test_parse_generic_code_workbook_multiline():
    source = OfficialSource("x", "近畿厚生局", "大阪", "https://kouseikyoku.mhlw.go.jp/x.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="大阪府", expected_prefectures=("大阪府",))
    frame = parse_code_workbook(make_book(), "大阪医科.xlsx", source)
    assert len(frame) == 4
    first = frame.iloc[0]
    assert first["clinic_id"] == "大阪府:0112345"
    assert first["phone"] == "06-1234-5678"
    assert first["facility_type"] == "診療所"
    assert first["status"] == "現存"
    assert first["registration_reason"] == "新規"
    assert first["as_of"] == "2026-09-01"


def test_population_filters_and_completed_ids():
    source = OfficialSource("x", "近畿厚生局", "大阪", "https://kouseikyoku.mhlw.go.jp/x.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="大阪府", expected_prefectures=("大阪府",))
    frame = parse_code_workbook(make_book(), "大阪医科.xlsx", source)
    targets, counts = build_maps_population(frame, {"大阪府:0112345"})
    assert targets == []
    assert counts["already_maps_completed"] == 1
    assert counts["hospital"] == 1
    assert counts["center"] == 1
    assert counts["inactive"] == 1


def test_three_way_split_exact_and_no_duplicates():
    records = [
        {"clinic_id": f"東京都:{i:07d}", "clinic_name": f"医院{i}", "phone": "03-0000-0000", "address": "東京都", "prefecture": "東京都", "medical_type": "医科", "facility_type": "診療所", "status": "現存"}
        for i in range(101)
    ]
    groups = split_three(records)
    assert sorted(len(g) for g in groups) == [33, 34, 34]
    keys = [r["clinic_id"] for g in groups for r in g]
    assert len(keys) == len(set(keys)) == 101
    assert [[r["clinic_id"] for r in g] for g in groups] == [[r["clinic_id"] for r in g] for g in split_three(records)]


def test_package_uses_exact_existing_queue_headers():
    rows = [{
        "clinic_id": f"神奈川県:{i:07d}", "clinic_name": f"医院{i}", "address": "神奈川県横浜市", "phone": "045-000-0000",
        "prefecture": "神奈川県", "medical_type": "医科", "facility_type": "診療所", "owner_name": "", "manager_name": "",
        "designation_date": "", "registration_reason": "", "status": "現存", "as_of": "2026-09-01", "source_url": "",
        "designation_period_start": "", "departments": "内", "source_bureau": "関東信越厚生局", "source_file": "x.xlsx",
    } for i in range(5)]
    frame = pd.DataFrame(rows, columns=MASTER_HEADERS)
    package, manifest, files = build_package(frame, set(), [], {"completed": 0})
    assert manifest["pc_counts"] == {"pc1": 2, "pc2": 2, "pc3": 1}
    with zipfile.ZipFile(BytesIO(package)) as zf:
        assert {"maps_queue_pc1.csv", "maps_queue_pc2.csv", "maps_queue_pc3.csv", "maps_queue_all.csv", "national_kouseikyoku_master.csv", "manifest.json", "README_FIRST.txt"}.issubset(zf.namelist())
        header = zf.read("maps_queue_pc1.csv").decode("utf-8-sig").splitlines()[0].split(",")
        assert header == QUEUE_HEADERS


def test_real_tokyo_workbook_if_present():
    path = Path("data/raw/131コード内容別一覧表（医科）東京r0809.xlsx")
    if not path.exists():
        return
    source = OfficialSource("tokyo", "関東信越厚生局", "東京", "https://kouseikyoku.mhlw.go.jp/x.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", expected_prefectures=("東京都",))
    frame = parse_code_workbook(path.read_bytes(), path.name, source)
    assert len(frame) > 10000
    assert set(frame["prefecture"]) == {"東京都"}
    assert set(frame["medical_type"]) == {"医科"}
    assert frame["clinic_id"].str.startswith("東京都:").all()
    assert frame["facility_type"].isin(["病院", "診療所"]).all()


def test_official_registry_covers_all_47_once():
    covered = [pref for source in OFFICIAL_SOURCES for pref in source.expected_prefectures]
    assert len(covered) == 47
    assert len(set(covered)) == 47
    assert set(covered) == set(PREFECTURES)
    assert all(source.url.startswith("https://kouseikyoku.mhlw.go.jp/") for source in OFFICIAL_SOURCES)


def test_parse_tohoku_hyphen_medical_code_and_primary_number_only():
    book = Workbook(); ws = book.active
    ws.title = "青森県"
    ws.append(["コード内容別医療機関一覧表"])
    ws.append(["[令和 8年 9月 1日現在　医科　現存/休止]"])
    ws.append(["1", "01-1024-3", "青森見本クリニック", "〒030-0000 青森市見本1-1", "017-123-4567", "A", "B", "令2.1.1\n新規\n令2.1.1", "内", "診療所\n現存"])
    ws.append(["2", "01-1216-5\n(01-3216-9)", "青森併設見本クリニック", "〒030-0001 青森市見本1-2", "017-123-4568", "A", "B", "令2.1.1", "内", "診療所\n現存"])
    out = BytesIO(); book.save(out)
    source = OfficialSource("tohoku_test", "東北厚生局", "青森県", "https://kouseikyoku.mhlw.go.jp/x.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="青森県", expected_prefectures=("青森県",))
    frame = parse_code_workbook(out.getvalue(), "青森県医科.xlsx", source)
    assert frame["clinic_id"].tolist() == ["青森県:0110243", "青森県:0112165"]


def test_parse_fullwidth_dash_medical_code():
    book = Workbook(); ws = book.active
    ws.title = "福島県"
    ws.append(["コード内容別医療機関一覧表"])
    ws.append(["[令和 8年 9月 1日現在　医科　現存/休止]"])
    ws.append(["1", "02－1340－6", "福島見本クリニック", "〒960-0000 福島市見本1-1", "024-123-4567", "A", "B", "令2.1.1", "内", "診療所\n現存"])
    out = BytesIO(); book.save(out)
    source = OfficialSource("tohoku_test2", "東北厚生局", "福島県", "https://kouseikyoku.mhlw.go.jp/x.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="福島県", expected_prefectures=("福島県",))
    frame = parse_code_workbook(out.getvalue(), "福島県医科.xlsx", source)
    assert frame.iloc[0]["clinic_id"] == "福島県:0213406"


def test_parse_tohoku_plain_digit_medical_code():
    book = Workbook(); ws = book.active
    ws.title = "宮城県"
    ws.append(["コード内容別医療機関一覧表"])
    ws.append(["[令和 8年 9月 1日現在　医科　現存/休止]"])
    ws.append(["1", "0211239", "宮城見本クリニック", "〒980-0000 仙台市見本1-1", "022-123-4567", "A", "B", "令2.1.1", "内", "診療所\n現存"])
    ws.append(["2", "0210017\n(0231209)", "宮城併設見本クリニック", "〒980-0001 仙台市見本1-2", "022-123-4568", "A", "B", "令2.1.1", "内", "診療所\n現存"])
    out = BytesIO(); book.save(out)
    source = OfficialSource("tohoku_test3", "東北厚生局", "宮城県", "https://kouseikyoku.mhlw.go.jp/x.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="宮城県", expected_prefectures=("宮城県",))
    frame = parse_code_workbook(out.getvalue(), "宮城県医科.xlsx", source)
    assert frame["clinic_id"].tolist() == ["宮城県:0211239", "宮城県:0210017"]


def _one_record_book(pref: str, code: str = "01,1234,5") -> bytes:
    book = Workbook(); ws = book.active
    ws.title = "Sheet1"
    ws.append(["コード内容別医療機関一覧表"])
    ws.append(["[令和 8年 9月 1日現在　医科　現存/休止]"])
    ws.append(["1", code, f"{pref}見本クリニック", "〒100-0001 市区町村見本1-1", "03-1234-5678", "A", "B", "令2.1.1", "内", "診療所\n現存"])
    out = BytesIO(); book.save(out); return out.getvalue()


def test_kanto_all_10_prefectures_from_jis_filename_prefixes():
    from src.master.national_maps_queue import parse_official_payload
    kanto = next(s for s in OFFICIAL_SOURCES if s.source_id == "kanto_shinetsu")
    jis = {"茨城県":"08", "栃木県":"09", "群馬県":"10", "埼玉県":"11", "千葉県":"12", "東京都":"13", "神奈川県":"14", "新潟県":"15", "山梨県":"19", "長野県":"20"}
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for pref, code2 in jis.items():
            # 実ファイルと同様に JIS 2桁 + 帳票区分1 + 日本語名。
            zf.writestr(f"{code2}1コード内容別一覧表（医科）{pref}r0809.xlsx", _one_record_book(pref))
    frame = parse_official_payload(out.getvalue(), kanto)
    assert set(frame["prefecture"]) == set(kanto.expected_prefectures)
    assert len(frame) == 10


def test_prefecture_can_be_recovered_when_japanese_zip_name_is_mojibake():
    from src.master.national_maps_queue import _prefecture_from_identifier, _repair_zip_filename
    original = "111コード内容別一覧表（医科）埼玉r0809.xlsx"
    mojibake = original.encode("cp932").decode("cp437")
    assert "埼玉" not in mojibake
    assert "埼玉" in _repair_zip_filename(mojibake)
    assert _prefecture_from_identifier(mojibake, ("埼玉県",)) == "埼玉県"


def test_official_medical_code_separator_variants_seen_across_regions():
    from src.master.national_maps_queue import _extract_medical_code
    samples = {
        "030,176,2": "0301762",   # 埼玉型
        "01-0198-0": "0101980",   # 千葉/東北型
        "010,006.5": "0100065",   # 神奈川型
        "011,120,4": "0111204",   # 新潟型
        "01,0020,2": "0100202",   # 群馬型
        "01・1206・4": "0112064", # 長野型
        "02－1340－6": "0213406", # 全角ハイフン型
        "0211239": "0211239",     # 区切りなし型
    }
    for raw, expected in samples.items():
        assert _extract_medical_code(raw) == expected, raw


def test_jis_prefix_inference_covers_all_47_without_japanese_names():
    from src.master.national_maps_queue import PREFECTURE_TO_JIS, _prefecture_from_identifier
    for pref in PREFECTURES:
        code = PREFECTURE_TO_JIS[pref]
        assert _prefecture_from_identifier(f"{code}1xxxxxxxx_r0809.xlsx", (pref,)) == pref


def test_collect_all_official_aggregates_source_errors(monkeypatch, tmp_path):
    import src.master.national_maps_queue as nmq
    good = OfficialSource("good", "A局", "良い県", "https://kouseikyoku.mhlw.go.jp/good.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="大阪府", expected_prefectures=("大阪府",))
    bad1 = OfficialSource("bad1", "B局", "失敗1", "https://kouseikyoku.mhlw.go.jp/bad1.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="京都府", expected_prefectures=("京都府",))
    bad2 = OfficialSource("bad2", "C局", "失敗2", "https://kouseikyoku.mhlw.go.jp/bad2.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="兵庫県", expected_prefectures=("兵庫県",))
    payloads = {"good": _one_record_book("大阪府")}
    def fake_download(source, cache_dir, force=False, timeout=45):
        if source.source_id in payloads:
            p = tmp_path / (source.source_id + ".xlsx"); p.write_bytes(payloads[source.source_id])
            return payloads[source.source_id], p, False
        raise RuntimeError("test download failure")
    monkeypatch.setattr(nmq, "download_source", fake_download)
    try:
        nmq.collect_all_official(tmp_path, sources=(good, bad1, bad2))
        assert False, "expected combined error"
    except ValueError as exc:
        msg = str(exc)
        assert "失敗1" in msg and "失敗2" in msg
        assert "全国15ソースを最後まで検証" in msg


def test_mojibake_zip_entry_is_parsed_end_to_end():
    from src.master.national_maps_queue import parse_official_payload
    source = OfficialSource(
        "saitama_zip", "関東信越厚生局", "埼玉県テスト",
        "https://kouseikyoku.mhlw.go.jp/x.zip", "https://kouseikyoku.mhlw.go.jp/", "zip",
        expected_prefectures=("埼玉県",),
    )
    original = "111コード内容別一覧表（医科）埼玉r0809.xlsx"
    mojibake = original.encode("cp932").decode("cp437")
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(mojibake, _one_record_book("埼玉県"))
    frame = parse_official_payload(out.getvalue(), source)
    assert len(frame) == 1
    assert frame.iloc[0]["prefecture"] == "埼玉県"
