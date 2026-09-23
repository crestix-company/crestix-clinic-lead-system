"""全国の地方厚生（支）局データから Google Maps Collector 用キューを作る。

重要:
- このモジュールは clinics.sqlite3 を更新しない。既存 Maps 完了済みIDの読取だけ行う。
- Google Maps Collector v5.7.0 が読む既存8列をそのまま出力する。
- 公式配布ファイルは staging にキャッシュし、元の厚生局データも master CSV に残す。
"""
from __future__ import annotations

from dataclasses import asdict, dataclass, replace
from datetime import datetime, timezone
from io import BytesIO, StringIO
from pathlib import Path
from urllib.parse import urlparse
import csv
import hashlib
import json
import re
import sqlite3
import time
import unicodedata
import zipfile

import pandas as pd
import requests
from openpyxl import load_workbook

from src.master.google_maps import QUEUE_HEADERS, excluded_reason
from src.utils.date_utils import parse_date


NATIONAL_MAPS_PARSER_VERSION = "1.2-all-regions"

PREFECTURES = (
    "北海道", "青森県", "岩手県", "宮城県", "秋田県", "山形県", "福島県",
    "茨城県", "栃木県", "群馬県", "埼玉県", "千葉県", "東京都", "神奈川県",
    "新潟県", "富山県", "石川県", "福井県", "山梨県", "長野県", "岐阜県",
    "静岡県", "愛知県", "三重県", "滋賀県", "京都府", "大阪府", "兵庫県",
    "奈良県", "和歌山県", "鳥取県", "島根県", "岡山県", "広島県", "山口県",
    "徳島県", "香川県", "愛媛県", "高知県", "福岡県", "佐賀県", "長崎県",
    "熊本県", "大分県", "宮崎県", "鹿児島県", "沖縄県",
)

# 総務省/JIS都道府県コード。厚生局ZIPのファイル名は
# 例: 111...埼玉, 121...千葉, 131...東京 のように先頭2桁が都道府県コード、
# 3桁目が帳票区分になっていることがある。日本語ファイル名が文字化けしても
# このコードから都道府県を安全に復元できる。
JIS_PREFECTURE_CODES = {f"{i:02d}": pref for i, pref in enumerate(PREFECTURES, 1)}
PREFECTURE_TO_JIS = {pref: code for code, pref in JIS_PREFECTURE_CODES.items()}

# 近畿など一部ZIPは都道府県名をローマ字で持つ（例: kikanzentai_hyogo_ika.xlsx）。
# 日本語名やJISコードだけに依存せず、47都道府県の一般的なローマ字名も復元に使う。
ROMAN_PREFECTURE_ALIASES = {
    "hokkaido": "北海道", "aomori": "青森県", "iwate": "岩手県", "miyagi": "宮城県",
    "akita": "秋田県", "yamagata": "山形県", "fukushima": "福島県", "ibaraki": "茨城県",
    "tochigi": "栃木県", "gunma": "群馬県", "saitama": "埼玉県", "chiba": "千葉県",
    "tokyo": "東京都", "kanagawa": "神奈川県", "niigata": "新潟県", "toyama": "富山県",
    "ishikawa": "石川県", "fukui": "福井県", "yamanashi": "山梨県", "nagano": "長野県",
    "gifu": "岐阜県", "shizuoka": "静岡県", "aichi": "愛知県", "mie": "三重県",
    "shiga": "滋賀県", "kyoto": "京都府", "osaka": "大阪府", "hyogo": "兵庫県",
    "nara": "奈良県", "wakayama": "和歌山県", "tottori": "鳥取県", "shimane": "島根県",
    "okayama": "岡山県", "hiroshima": "広島県", "yamaguchi": "山口県", "tokushima": "徳島県",
    "kagawa": "香川県", "ehime": "愛媛県", "kochi": "高知県", "fukuoka": "福岡県",
    "saga": "佐賀県", "nagasaki": "長崎県", "kumamoto": "熊本県", "oita": "大分県",
    "miyazaki": "宮崎県", "kagoshima": "鹿児島県", "okinawa": "沖縄県",
}

# ZIPの古い日本語ファイル名はUTF-8フラグ無しでShift-JIS/CP932が使われることがあり、
# Python zipfile上ではCP437文字列として見える。その場合だけ可逆変換を試す。
def _repair_zip_filename(name: str) -> str:
    # 重要: CP932の2バイト目に0x5cが現れることがあるため、文字化け復元前に
    # バックスラッシュをパス区切りへ置換してはいけない。
    raw = str(name or "")
    candidates = [raw]
    try:
        repaired = raw.encode("cp437").decode("cp932")
        candidates.append(repaired)
    except (UnicodeEncodeError, UnicodeDecodeError):
        pass

    def score(text: str) -> tuple[int, int]:
        useful = sum(1 for token in ("コード", "内容", "一覧", "医科", "病院", "診療所") if token in text)
        useful += sum(4 for pref in PREFECTURES if pref in text)
        japanese = sum(1 for ch in text if "\u3040" <= ch <= "\u30ff" or "\u4e00" <= ch <= "\u9fff")
        return useful, japanese

    chosen = max(candidates, key=score)
    return chosen.replace("\\", "/")


def _prefecture_from_identifier(text: str, expected_prefectures=()) -> str:
    """ファイル名/シート名から都道府県を復元する。

    日本語名 → JIS都道府県コードの順で判定する。数字判定はbasename先頭だけに
    制限し、r0809等の日付を県コードと誤認しない。
    """
    repaired = _repair_zip_filename(text)
    pref = _prefecture_in_text(repaired, allow_short=True)
    expected = set(expected_prefectures or ())
    if pref and (not expected or pref in expected):
        return pref

    # 例: 2026.9_kikanzentai_hyogo_ika.xlsx。年をJISコードと誤認する前に
    # ローマ字都道府県名を優先して確定する。
    lower = repaired.lower()
    for alias, candidate in ROMAN_PREFECTURE_ALIASES.items():
        if re.search(rf"(?<![a-z]){re.escape(alias)}(?![a-z])", lower):
            if not expected or candidate in expected:
                return candidate

    base = Path(repaired.replace("\\", "/")).name
    # 典型: 111コード... = JIS 11(埼玉) + 帳票区分1。
    m = re.match(r"^[^0-9]{0,8}([0-4][0-9])(?:1)?(?=[^0-9]|$)", base)
    if not m:
        # 先頭に3桁以上続く場合も先頭2桁をJISとして使う（131..., 271...等）。
        m2 = re.match(r"^[^0-9]{0,8}([0-4][0-9])(?=[0-9])", base)
        code = m2.group(1) if m2 else ""
    else:
        code = m.group(1)
    candidate = JIS_PREFECTURE_CODES.get(code, "")
    if candidate and (not expected or candidate in expected):
        return candidate
    return ""

MASTER_HEADERS = [
    "clinic_id", "clinic_name", "address", "phone", "prefecture", "medical_type",
    "facility_type", "owner_name", "manager_name", "designation_date",
    "registration_reason", "status", "as_of", "source_url",
    "designation_period_start", "departments", "source_bureau", "source_file",
]

FINAL_MAPS_STATUSES = {
    "MAPS_MATCHED_WEBSITE", "MAPS_MATCHED_NO_WEBSITE", "MAPS_NOT_FOUND",
    "EXCLUDED_HOSPITAL", "EXCLUDED_CENTER",
}
BAD_OPERATION_STATUSES = ("休止", "廃止", "辞退", "取消")
# 医療機関番号は地方により表記が異なるが、主番号は最終的に7桁。
# 例: 01,1004,9 / 01-1024-3 / 01-15202 / 011 026.5 / 0211239
MEDICAL_CODE_7DIGIT_RE = re.compile(
    r"(?<!\d)((?:\d[\s,，・･\.．\-‐‑‒–—−ー]?){6}\d)(?!\d)"
)
DATE_CONTEXT_RE = re.compile(r"((?:令和|平成|昭和)\s*\d+年\s*\d+月\s*\d+日)\s*現在")


@dataclass(frozen=True)
class OfficialSource:
    source_id: str
    bureau: str
    label: str
    url: str
    index_url: str
    kind: str  # xlsx / zip
    expected_as_of: str = "2026-09-01"
    prefecture_hint: str = ""
    facility_type_hint: str = ""
    expected_prefectures: tuple[str, ...] = ()


OFFICIAL_SOURCES = (
    OfficialSource(
        "hokkaido_clinic", "北海道厚生局", "北海道・医科（診療所）", 
        "https://kouseikyoku.mhlw.go.jp/hokkaido/000499339.xlsx",
        "https://kouseikyoku.mhlw.go.jp/hokkaido/gyomu/gyomu/hoken_kikan/code_ichiran.html",
        "xlsx", prefecture_hint="北海道", facility_type_hint="診療所", expected_prefectures=("北海道",),
    ),
    OfficialSource(
        "tohoku", "東北厚生局", "東北6県・医科",
        "https://kouseikyoku.mhlw.go.jp/tohoku/shitei-touhoku-ika-r0809.xlsx",
        "https://kouseikyoku.mhlw.go.jp/tohoku/gyomu/gyomu/hoken_kikan/itiran.html",
        "xlsx", expected_prefectures=("青森県", "岩手県", "宮城県", "秋田県", "山形県", "福島県"),
    ),
    OfficialSource(
        "kanto_shinetsu", "関東信越厚生局", "関東信越10都県・医科",
        "https://kouseikyoku.mhlw.go.jp/kantoshinetsu/shitei_ika_r0809.zip",
        "https://kouseikyoku.mhlw.go.jp/kantoshinetsu/chousa/shitei.html",
        "zip", expected_prefectures=("茨城県", "栃木県", "群馬県", "埼玉県", "千葉県", "東京都", "神奈川県", "新潟県", "山梨県", "長野県"),
    ),
    OfficialSource(
        "tokai_hokuriku", "東海北陸厚生局", "東海北陸6県・医科",
        "https://kouseikyoku.mhlw.go.jp/tokaihokuriku/2609-01-01.zip",
        "https://kouseikyoku.mhlw.go.jp/tokaihokuriku/newpage_00287.html",
        "zip", expected_prefectures=("富山県", "石川県", "岐阜県", "静岡県", "愛知県", "三重県"),
    ),
    OfficialSource(
        "kinki", "近畿厚生局", "近畿7府県・医科",
        "https://kouseikyoku.mhlw.go.jp/kinki/2026.9_kikanzentai_ika.zip",
        "https://kouseikyoku.mhlw.go.jp/kinki/tyousa/shinkishitei.html",
        "zip", expected_prefectures=("福井県", "滋賀県", "京都府", "大阪府", "兵庫県", "奈良県", "和歌山県"),
    ),
    OfficialSource(
        "chugoku", "中国四国厚生局", "中国5県・医科",
        "https://kouseikyoku.mhlw.go.jp/chugokushikoku/000500018.zip",
        "https://kouseikyoku.mhlw.go.jp/chugokushikoku/chousaka/iryoukikanshitei.html",
        "zip", expected_prefectures=("鳥取県", "島根県", "岡山県", "広島県", "山口県"),
    ),
    OfficialSource(
        "shikoku", "四国厚生支局", "四国4県・医科",
        "https://kouseikyoku.mhlw.go.jp/shikoku/000499483.zip",
        "https://kouseikyoku.mhlw.go.jp/shikoku/gyomu/gyomu/hoken_kikan/shitei/",
        "zip", expected_prefectures=("徳島県", "香川県", "愛媛県", "高知県"),
    ),
    OfficialSource(
        "kyushu_fukuoka", "九州厚生局", "福岡県・エクセルデータ",
        "https://kouseikyoku.mhlw.go.jp/kyushu/000500310.zip",
        "https://kouseikyoku.mhlw.go.jp/kyushu/gyomu/gyomu/hoken_kikan/index_00006.html",
        "zip", prefecture_hint="福岡県", expected_prefectures=("福岡県",),
    ),
    OfficialSource(
        "kyushu_saga", "九州厚生局", "佐賀県・エクセルデータ",
        "https://kouseikyoku.mhlw.go.jp/kyushu/000500311.zip",
        "https://kouseikyoku.mhlw.go.jp/kyushu/gyomu/gyomu/hoken_kikan/index_00006.html",
        "zip", prefecture_hint="佐賀県", expected_prefectures=("佐賀県",),
    ),
    OfficialSource(
        "kyushu_nagasaki", "九州厚生局", "長崎県・エクセルデータ",
        "https://kouseikyoku.mhlw.go.jp/kyushu/000500315.zip",
        "https://kouseikyoku.mhlw.go.jp/kyushu/gyomu/gyomu/hoken_kikan/index_00006.html",
        "zip", prefecture_hint="長崎県", expected_prefectures=("長崎県",),
    ),
    OfficialSource(
        "kyushu_kumamoto", "九州厚生局", "熊本県・エクセルデータ",
        "https://kouseikyoku.mhlw.go.jp/kyushu/000500316.zip",
        "https://kouseikyoku.mhlw.go.jp/kyushu/gyomu/gyomu/hoken_kikan/index_00006.html",
        "zip", prefecture_hint="熊本県", expected_prefectures=("熊本県",),
    ),
    OfficialSource(
        "kyushu_oita", "九州厚生局", "大分県・エクセルデータ",
        "https://kouseikyoku.mhlw.go.jp/kyushu/000500317.zip",
        "https://kouseikyoku.mhlw.go.jp/kyushu/gyomu/gyomu/hoken_kikan/index_00006.html",
        "zip", prefecture_hint="大分県", expected_prefectures=("大分県",),
    ),
    OfficialSource(
        "kyushu_miyazaki", "九州厚生局", "宮崎県・エクセルデータ",
        "https://kouseikyoku.mhlw.go.jp/kyushu/000500318.zip",
        "https://kouseikyoku.mhlw.go.jp/kyushu/gyomu/gyomu/hoken_kikan/index_00006.html",
        "zip", prefecture_hint="宮崎県", expected_prefectures=("宮崎県",),
    ),
    OfficialSource(
        "kyushu_kagoshima", "九州厚生局", "鹿児島県・エクセルデータ",
        "https://kouseikyoku.mhlw.go.jp/kyushu/000500321.zip",
        "https://kouseikyoku.mhlw.go.jp/kyushu/gyomu/gyomu/hoken_kikan/index_00006.html",
        "zip", prefecture_hint="鹿児島県", expected_prefectures=("鹿児島県",),
    ),
    OfficialSource(
        "kyushu_okinawa", "九州厚生局", "沖縄県・エクセルデータ",
        "https://kouseikyoku.mhlw.go.jp/kyushu/000500320.zip",
        "https://kouseikyoku.mhlw.go.jp/kyushu/gyomu/gyomu/hoken_kikan/index_00006.html",
        "zip", prefecture_hint="沖縄県", expected_prefectures=("沖縄県",),
    ),
)


def _s(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and pd.isna(value):
        return ""
    return unicodedata.normalize("NFKC", str(value)).strip()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _csv_bytes(headers, rows) -> bytes:
    out = StringIO(newline="")
    writer = csv.writer(out, lineterminator="\r\n")
    writer.writerow(headers)
    writer.writerows(rows)
    return out.getvalue().encode("utf-8-sig")


def _normalize_medical_code(value: str) -> str:
    return re.sub(r"\D", "", _s(value))


def _extract_medical_code(value: str) -> str:
    """医療機関番号セルから主番号7桁だけを抽出する。

    地方差を区切り文字の個数で決め打ちしない。近畿の ``01-15202``、四国の
    ``011 026.5``、関東の ``01,1004,9``、東北の ``01-1024-3``、区切りなし
    ``0211239`` を同じ7桁として扱う。同一セルに併設番号が続く場合は最初の
    7桁を主番号とする。
    """
    text = _s(value).replace("\r", "\n")
    if not text:
        return ""

    # まず改行・括弧単位で「数字だけにすると7桁」の主番号を優先する。
    # これで ``01-01813\n(01-61813)`` の先頭番号だけを安定して採用できる。
    parts = re.split(r"[\n()（）\[\]［］]+", text)
    for part in parts:
        digits = re.sub(r"\D", "", part)
        if len(digits) == 7:
            return digits

    # ラベル文字が同じ行に付く地域向け。7数字の並びを最初に見つける。
    match = MEDICAL_CODE_7DIGIT_RE.search(text)
    if match:
        digits = re.sub(r"\D", "", match.group(1))
        if len(digits) == 7:
            return digits
    return ""


def _find_medical_code_cell(cells: list[str]) -> tuple[int, str]:
    """帳票先頭付近から医療機関番号列を探す。通常は2列目だが地方差を許容。"""
    order = [1, 0, 2, 3]
    for idx in order:
        if idx >= len(cells):
            continue
        code = _extract_medical_code(cells[idx])
        if code:
            return idx, code
    return -1, ""


def _prefecture_in_text(text: str, allow_short: bool = False) -> str:
    normalized = _s(text)
    # 住所などでは正式都道府県名だけを見る。ファイル名/シート名では「東京」「大阪」等も許容。
    for pref in PREFECTURES:
        if pref in normalized:
            return pref
    if allow_short:
        aliases = sorted(
            ((pref[:-1] if pref.endswith(("都", "府", "県")) else pref, pref) for pref in PREFECTURES),
            key=lambda x: len(x[0]), reverse=True,
        )
        for alias, pref in aliases:
            if alias and alias in normalized:
                return pref
    return ""


def _clean_multiline(value: str) -> str:
    parts = [re.sub(r"\s+", " ", p).strip() for p in _s(value).replace("\r", "\n").split("\n")]
    return " ".join(p for p in parts if p)


def _clean_address(value: str, prefecture: str) -> str:
    text = _clean_multiline(value)
    text = re.sub(r"^〒\s*[0-9０-９]{3}\s*[-−ー―‐]?\s*[0-9０-９]{4}\s*", "", text)
    text = text.strip()
    if prefecture and text and not text.startswith(prefecture):
        text = prefecture + text
    return text


def _extract_phone(value: str) -> str:
    text = _s(value).replace("\r", "\n")
    first = next((x.strip() for x in text.split("\n") if x.strip()), "")
    # Excelによっては同一セルに勤務医数が続く。最初の電話らしい部分だけを採用。
    match = re.search(r"(?:\(?0\d{1,4}\)?[-‐‑‒–—−ー\s]?)?\d{1,4}[-‐‑‒–—−ー\s]\d{3,4}[-‐‑‒–—−ー\s]\d{3,4}|0\d{8,10}|0\d{1,4}\(\d{1,4}\)\d{3,4}", first)
    return match.group(0).strip() if match else first.split()[0] if first else ""


def _department_text(value: str) -> str:
    text = _clean_multiline(value)
    text = re.sub(r"(?:一般|療養|精神|感染|結核)\s*[0-9０-９]+", "", text)
    text = re.sub(r"\s+", " ", text).strip()
    return "" if re.fullmatch(r"[\s0-9０-９]*", text) else text


def _parse_designation(value: str) -> tuple[str, str, str]:
    raw = _s(value).replace("\r", "\n")
    parts = [p.strip() for p in raw.split("\n") if p.strip()]
    dates: list[str] = []
    reason = ""
    for part in parts:
        parsed = parse_date(part)
        if parsed:
            dates.append(parsed.isoformat())
        elif not reason and len(part) <= 40:
            reason = part
    designation = dates[0] if dates else ""
    period_start = dates[1] if len(dates) > 1 else ""
    if not designation:
        parsed = parse_date(raw)
        designation = parsed.isoformat() if parsed else raw
    return designation, reason, period_start


def _facility_and_status(value: str, hint: str = "") -> tuple[str, str]:
    text = _clean_multiline(value)
    facility = hint if hint in ("病院", "診療所") else ""
    if "診療所" in text:
        facility = "診療所"
    elif "病院" in text or "特定機能" in text:
        facility = "病院"
    status = ""
    for candidate in ("現存", "休止", "廃止", "辞退", "取消"):
        if candidate in text:
            status = candidate
            break
    return facility, status


def _max_reported_beds(value: str) -> int:
    """病床数/診療科欄から明示された病床数の最大値を返す。

    病院は20床以上、診療所は19床以下という法的な区分を、備考欄の文字が
    地方差で別セルへずれた場合の安全な補助判定にだけ使う。
    """
    text = unicodedata.normalize("NFKC", _clean_multiline(value))
    counts: list[int] = []
    for marker in ("一般", "療養", "精神", "感染", "結核"):
        for match in re.finditer(rf"{marker}[^0-9]{{0,8}}([0-9]{{1,4}})", text):
            try:
                counts.append(int(match.group(1)))
            except ValueError:
                pass
    return max(counts) if counts else 0


def _infer_facility_and_status(
    clinic_name: str, evidence: str, hint: str = "", existing_facility: str = "", existing_status: str = ""
) -> tuple[str, str]:
    """帳票末尾の備考/病床数/診療科欄をまとめて病院・診療所を確定する。

    地域ごとのExcelでは「地域支援」「病院」「現存」が別セルに分かれることがある。
    1セル決め打ちをやめ、末尾ブロック全体で判定する。最後まで明示区分が無い場合は
    医療機関名と病床数を安全ガードとして使う。
    """
    facility, status = _facility_and_status(evidence, existing_facility or hint)
    if not status:
        status = existing_status
    if facility not in ("病院", "診療所"):
        name = unicodedata.normalize("NFKC", _s(clinic_name))
        if "病院" in name:
            facility = "病院"
        elif _max_reported_beds(evidence) >= 20:
            facility = "病院"
        else:
            # 「センター」は後段の excluded_reason() で必ず除外される。
            # それ以外で病院の強い根拠が無い医科一覧の行は診療所として扱う。
            facility = "診療所"
    return facility, status

def _sheet_preview(ws, max_rows: int = 30) -> str:
    values: list[str] = [ws.title]
    for row in ws.iter_rows(min_row=1, max_row=min(max_rows, ws.max_row), values_only=True):
        values.extend(_s(v) for v in row if _s(v))
    return " ".join(values)


def _looks_non_medical(filename: str, preview: str) -> bool:
    text = _s(filename + " " + preview).lower()
    # 基準行が最優先。医科と明示されていれば医院名中の「歯科」等では除外しない。
    basis = re.search(r"(?:現在|コード内容別医療機関一覧表).{0,80}", text)
    target = basis.group(0) if basis else text[:1000]
    if "薬局" in target or "薬 科" in target:
        return True
    if "歯科" in target and "医科" not in target:
        return True
    lower_name = _s(filename).lower()
    if any(k in lower_name for k in ("yaku", "yakkyoku", "薬局", "shika", "sika", "歯科")) and not any(k in lower_name for k in ("ika", "医科")):
        return True
    return False


def parse_code_workbook(raw: bytes, filename: str, source: OfficialSource) -> pd.DataFrame:
    """標準的なコード内容別医療機関一覧表を全国共通形式へ変換する。

    県別ファイル、県別シート、1シート中の県切替に対応する。
    地方差が大きいのは主に医療機関番号表記と備考欄のセル位置なので、
    それらを1形式に決め打ちせず、主番号7桁と末尾ブロック全体から復元する。
    """
    book = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    all_records: list[dict] = []
    try:
        for ws in book.worksheets:
            preview = _sheet_preview(ws)
            if _looks_non_medical(filename, preview):
                continue

            basis = ""
            match = DATE_CONTEXT_RE.search(preview)
            if match:
                d = parse_date(match.group(1))
                basis = d.isoformat() if d else ""
            if not basis:
                basis = source.expected_as_of

            sheet_pref = (
                source.prefecture_hint
                or _prefecture_from_identifier(filename, source.expected_prefectures)
                or _prefecture_from_identifier(ws.title, source.expected_prefectures)
                or _prefecture_in_text(filename + " " + ws.title, allow_short=True)
            )
            current_pref = sheet_pref
            record: dict | None = None

            def commit():
                nonlocal record
                if not record:
                    return
                evidence = _s(record.pop("_facility_evidence", ""))
                facility, status = _infer_facility_and_status(
                    record.get("clinic_name", ""), evidence, source.facility_type_hint,
                    record.get("facility_type", ""), record.get("status", ""),
                )
                record["facility_type"] = facility
                if status:
                    record["status"] = status
                record.pop("_code_idx", None)
                all_records.append(record)
                record = None

            for raw_row in ws.iter_rows(values_only=True):
                cells = [_s(v) for v in raw_row]
                cells += [""] * max(0, 14 - len(cells))
                row_text = " ".join(x for x in cells if x)
                nonempty = sum(bool(x) for x in cells)
                row_pref = _prefecture_in_text(row_text, allow_short=nonempty <= 3)

                code_idx, primary_code = _find_medical_code_cell(cells)
                serial_ok = False
                if code_idx > 0:
                    left = cells[max(0, code_idx - 2):code_idx]
                    serial_ok = any(x.isdigit() for x in left)
                elif code_idx == 0:
                    serial_ok = bool(cells[1])
                is_start = bool(primary_code) and serial_ok

                # 県見出しの行では現在県を切り替える。施設行の住所内都道府県も後で優先する。
                if row_pref and not is_start:
                    current_pref = row_pref
                if "現在" in row_text and not basis:
                    m = DATE_CONTEXT_RE.search(row_text)
                    if m:
                        d = parse_date(m.group(1))
                        basis = d.isoformat() if d else source.expected_as_of

                if is_start:
                    commit()
                    name_idx = code_idx + 1
                    address_idx = code_idx + 2
                    phone_idx = code_idx + 3
                    owner_idx = code_idx + 4
                    manager_idx = code_idx + 5
                    designation_idx = code_idx + 6
                    department_idx = code_idx + 7
                    tail_idx = department_idx

                    addr_pref = _prefecture_in_text(cells[address_idx])
                    pref = addr_pref or current_pref or sheet_pref
                    if not pref and len(source.expected_prefectures) == 1:
                        pref = source.expected_prefectures[0]

                    evidence = " ".join(x for x in cells[tail_idx:] if x)
                    facility, status = _facility_and_status(evidence, source.facility_type_hint)
                    designation, reason, period_start = _parse_designation(cells[designation_idx])
                    record = {
                        "clinic_id": f"{pref}:{primary_code}" if pref and primary_code else "",
                        "clinic_name": _clean_multiline(cells[name_idx]),
                        "address": _clean_address(cells[address_idx], pref),
                        "phone": _extract_phone(cells[phone_idx]),
                        "prefecture": pref,
                        "medical_type": "医科",
                        "facility_type": facility,
                        "owner_name": _clean_multiline(cells[owner_idx]),
                        "manager_name": _clean_multiline(cells[manager_idx]),
                        "designation_date": designation,
                        "registration_reason": reason,
                        "status": status,
                        "as_of": basis or source.expected_as_of,
                        "source_url": source.url,
                        "designation_period_start": period_start,
                        "departments": _department_text(cells[department_idx]),
                        "source_bureau": source.bureau,
                        "source_file": filename + (f"#{ws.title}" if len(book.worksheets) > 1 else ""),
                        "_facility_evidence": evidence,
                        "_code_idx": code_idx,
                    }
                    if pref:
                        current_pref = pref

                elif record and code_idx < 0 and not cells[0] and not cells[1]:
                    code_col = int(record.get("_code_idx", 1))
                    designation_idx = code_col + 6
                    department_idx = code_col + 7
                    tail_idx = department_idx

                    dep = _department_text(cells[department_idx])
                    if dep:
                        record["departments"] = (record["departments"] + " " + dep).strip()
                    evidence = " ".join(x for x in cells[tail_idx:] if x)
                    if evidence:
                        record["_facility_evidence"] = (record.get("_facility_evidence", "") + " " + evidence).strip()
                        facility, status = _facility_and_status(
                            evidence, record.get("facility_type") or source.facility_type_hint
                        )
                        if facility:
                            record["facility_type"] = facility
                        if status:
                            record["status"] = status

                    if cells[designation_idx]:
                        d, reason, period_start = _parse_designation(cells[designation_idx])
                        if reason and not record["registration_reason"]:
                            record["registration_reason"] = reason
                        if period_start:
                            record["designation_period_start"] = period_start
                        elif d and not record["designation_period_start"] and d != record["designation_date"]:
                            record["designation_period_start"] = d
            commit()

    finally:
        book.close()

    if not all_records:
        raise ValueError(f"医科のコード内容別医療機関一覧を読み取れませんでした: {filename}")
    frame = pd.DataFrame(all_records).fillna("")
    missing_pref = frame[frame["prefecture"] == ""]
    if not missing_pref.empty:
        sample = ", ".join(missing_pref["clinic_name"].head(3).tolist())
        raise ValueError(f"都道府県を特定できない行があります: {filename} / {sample}")

    # v1.2では区分1セル欠落だけで県全体を捨てない。上の安全推定で必ず二値化する。
    bad_facility = frame[~frame["facility_type"].isin(["病院", "診療所"])]
    if not bad_facility.empty:
        sample = ", ".join(bad_facility["clinic_name"].head(3).tolist())
        raise ValueError(f"病院/診療所区分を最終確定できない行があります: {filename} / {sample}")
    return frame[MASTER_HEADERS].copy()

def _xlsx_entries(payload: bytes, source: OfficialSource, source_filename: str) -> list[tuple[str, bytes]]:
    if source.kind == "xlsx":
        return [(source_filename, payload)]
    if not zipfile.is_zipfile(BytesIO(payload)):
        raise ValueError(f"公式ZIPとして読めません: {source.label}")
    results: list[tuple[str, bytes]] = []
    with zipfile.ZipFile(BytesIO(payload)) as zf:
        for info in zf.infolist():
            # 文字化け復元より前に\を置換するとCP932の2バイト文字を壊すため、
            # info.filenameの生文字列をそのまま復元関数へ渡す。
            raw_name = info.filename
            name = _repair_zip_filename(raw_name)
            safe_raw = raw_name.replace("\\", "/")
            if info.is_dir() or not name.lower().endswith(".xlsx"):
                continue
            if safe_raw.startswith("/") or "../" in safe_raw or info.file_size > 50 * 1024 * 1024:
                raise ValueError(f"ZIP内ファイルが安全条件を満たしません: {name}")
            results.append((Path(name).name, zf.read(info)))
    if not results:
        raise ValueError(f"ZIP内にExcelがありません: {source.label}")
    return results


def parse_official_payload(payload: bytes, source: OfficialSource) -> pd.DataFrame:
    suffix = ".xlsx" if source.kind == "xlsx" else ".zip"
    source_filename = source.source_id + suffix
    frames: list[pd.DataFrame] = []
    errors: list[str] = []
    entries = _xlsx_entries(payload, source, source_filename)
    for filename, raw in entries:
        # ZIP内の日本語名が文字化けしても、先頭JIS都道府県コードから県を確定する。
        entry_pref = source.prefecture_hint or _prefecture_from_identifier(filename, source.expected_prefectures)
        entry_source = replace(source, prefecture_hint=entry_pref) if entry_pref else source
        try:
            frame = parse_code_workbook(raw, filename, entry_source)
        except ValueError as exc:
            # 九州ZIP等には医科以外が同梱される場合があるため対象外Workbookは許容。
            # ただし最終的な都道府県カバレッジで必ず完全性を検証する。
            errors.append(f"{filename}: {exc}")
            continue
        frames.append(frame)
    if not frames:
        detail = " / ".join(errors[:5])
        raise ValueError(f"{source.label} の医科Excelを解析できませんでした。{detail}")
    result = pd.concat(frames, ignore_index=True).fillna("")
    result = result.drop_duplicates(subset=["clinic_id"], keep="last").reset_index(drop=True)
    present = set(result["prefecture"].unique())
    missing = set(source.expected_prefectures) - present
    extra = present - set(source.expected_prefectures)
    if missing:
        detail = " / ".join(errors[:3])
        suffix_detail = f" / 解析メモ: {detail}" if detail else ""
        raise ValueError(f"{source.label}: 都道府県が不足しています: {', '.join(sorted(missing))}{suffix_detail}")
    if extra and source.expected_prefectures:
        raise ValueError(f"{source.label}: 想定外の都道府県を検出しました: {', '.join(sorted(extra))}")
    return result[MASTER_HEADERS].copy()


def download_source(source: OfficialSource, cache_dir: Path, timeout: int = 90, force: bool = False) -> tuple[bytes, Path, bool]:
    """公式ドメインだけから取得し、キャッシュする。戻り値: bytes, path, cache_hit。"""
    cache_dir = Path(cache_dir)
    cache_dir.mkdir(parents=True, exist_ok=True)
    suffix = ".xlsx" if source.kind == "xlsx" else ".zip"
    target = cache_dir / f"{source.source_id}{suffix}"
    if target.exists() and target.stat().st_size > 1000 and not force:
        return target.read_bytes(), target, True
    parsed = urlparse(source.url)
    if parsed.scheme != "https" or parsed.hostname != "kouseikyoku.mhlw.go.jp":
        raise ValueError("公式厚生局ドメイン以外は自動取得しません。")
    session = requests.Session()
    session.headers["User-Agent"] = "ClinicListFilter/2.5.5 national-maps-precollection (public-data import)"
    last_error: Exception | None = None
    for attempt in range(3):
        try:
            response = session.get(source.url, timeout=timeout, allow_redirects=True)
            response.raise_for_status()
            final = urlparse(response.url)
            if final.scheme != "https" or final.hostname != "kouseikyoku.mhlw.go.jp":
                raise ValueError("公式厚生局ドメイン外へリダイレクトされたため停止しました。")
            payload = response.content
            if len(payload) > 80 * 1024 * 1024:
                raise ValueError("公式配布ファイルが80MBを超えたため停止しました。")
            if not payload.startswith(b"PK"):
                raise ValueError("Excel/ZIPではない応答を受け取ったため停止しました。")
            # xlsxもZIPコンテナなのでPKシグネチャで確認できる。
            target.write_bytes(payload)
            return payload, target, False
        except Exception as exc:  # requests系と安全検証を同じ再試行単位にする
            last_error = exc
            if attempt < 2:
                time.sleep(1.0 + attempt)
    raise RuntimeError(f"{source.label} のダウンロードに失敗しました: {last_error}") from last_error


def load_completed_maps_ids_readonly(db_path: Path) -> tuple[set[str], dict[str, int]]:
    """本番DBを read-only URI で読む。書込み・schema初期化はしない。"""
    db_path = Path(db_path).resolve()
    if not db_path.exists():
        return set(), {"rows": 0, "completed": 0}
    uri = "file:" + db_path.as_posix() + "?mode=ro"
    completed: set[str] = set()
    rows = 0
    conn = sqlite3.connect(uri, uri=True, timeout=10)
    try:
        conn.row_factory = sqlite3.Row
        for row in conn.execute(
            "SELECT effective_json,maps_presence_status FROM clinics WHERE merged_into IS NULL AND maps_presence_status<>''"
        ):
            rows += 1
            if _s(row["maps_presence_status"]) not in FINAL_MAPS_STATUSES:
                continue
            try:
                data = json.loads(row["effective_json"] or "{}")
            except json.JSONDecodeError:
                continue
            key = _s(data.get("clinic_id") or data.get("medical_institution_number"))
            if key:
                completed.add(key)
    finally:
        conn.close()
    return completed, {"rows": rows, "completed": len(completed)}


def _stable_key(record: dict) -> str:
    key = _s(record.get("clinic_id"))
    if key:
        return key
    return "|".join(_s(record.get(k)) for k in ("prefecture", "clinic_name", "phone", "address"))


def build_maps_population(frame: pd.DataFrame, completed_ids: set[str] | None = None) -> tuple[list[dict], dict]:
    completed_ids = completed_ids or set()
    counts = {
        "input": 0, "duplicates": 0, "non_medical": 0, "hospital": 0, "center": 0,
        "inactive": 0, "missing_identity": 0, "already_maps_completed": 0, "target": 0,
    }
    targets: list[dict] = []
    seen: set[str] = set()
    for raw in frame.fillna("").to_dict("records"):
        counts["input"] += 1
        record = {k: _s(v) for k, v in raw.items()}
        key = _stable_key(record)
        if not key or key in seen:
            counts["duplicates"] += 1
            continue
        seen.add(key)
        if record.get("medical_type") != "医科":
            counts["non_medical"] += 1
            continue
        reason = excluded_reason(record)
        if reason == "hospital" or record.get("facility_type") != "診療所":
            counts["hospital"] += 1
            continue
        if reason == "center":
            counts["center"] += 1
            continue
        status = record.get("status", "")
        if any(word in status for word in BAD_OPERATION_STATUSES):
            counts["inactive"] += 1
            continue
        if not record.get("clinic_name") or not (record.get("phone") or record.get("address")):
            counts["missing_identity"] += 1
            continue
        if record.get("clinic_id") in completed_ids:
            counts["already_maps_completed"] += 1
            continue
        targets.append(record)
    counts["target"] = len(targets)
    return targets, counts


def split_three(records: list[dict]) -> list[list[dict]]:
    """地理偏りを避けつつ再現可能に3等分。件数差は最大1。"""
    ordered = sorted(records, key=lambda r: hashlib.sha256(_stable_key(r).encode("utf-8")).hexdigest())
    groups = [[], [], []]
    for i, record in enumerate(ordered):
        groups[i % 3].append(record)
    return groups


def queue_bytes(records: list[dict]) -> bytes:
    rows = []
    for r in records:
        rows.append([
            "", r.get("clinic_id", ""), r.get("clinic_name", ""), r.get("phone", ""),
            r.get("address", ""), r.get("prefecture", ""), r.get("facility_type", ""), r.get("status", ""),
        ])
    return _csv_bytes(QUEUE_HEADERS, rows)


def master_bytes(frame: pd.DataFrame) -> bytes:
    rows = [[_s(row.get(h, "")) for h in MASTER_HEADERS] for row in frame.to_dict("records")]
    return _csv_bytes(MASTER_HEADERS, rows)


def build_package(frame: pd.DataFrame, completed_ids: set[str], source_report: list[dict], db_info: dict | None = None) -> tuple[bytes, dict, dict[str, bytes]]:
    targets, filters = build_maps_population(frame, completed_ids)
    groups = split_three(targets)
    files: dict[str, bytes] = {
        "maps_queue_all.csv": queue_bytes(targets),
        "maps_queue_pc1.csv": queue_bytes(groups[0]),
        "maps_queue_pc2.csv": queue_bytes(groups[1]),
        "maps_queue_pc3.csv": queue_bytes(groups[2]),
        "national_kouseikyoku_master.csv": master_bytes(frame),
    }
    per_pref = pd.Series([r.get("prefecture", "") for r in targets]).value_counts().sort_index().to_dict() if targets else {}
    manifest = {
        "created_at": datetime.now(timezone.utc).isoformat(),
        "purpose": "Google Maps Clinic Collector v5.7.0 precollection / 3 PCs",
        "database_writes": False,
        "queue_headers": QUEUE_HEADERS,
        "source_rows": int(len(frame)),
        "target_rows": len(targets),
        "pc_counts": {f"pc{i+1}": len(group) for i, group in enumerate(groups)},
        "per_prefecture_target": per_pref,
        "filter_counts": filters,
        "existing_db_readonly": db_info or {},
        "sources": source_report,
    }
    file_hashes = {name: hashlib.sha256(data).hexdigest() for name, data in files.items()}
    manifest["sha256"] = file_hashes
    readme = """全国 Google Maps 3台並列収集パッケージ\n\n1. PC1には maps_queue_pc1.csv、PC2には maps_queue_pc2.csv、PC3には maps_queue_pc3.csv を渡します。\n2. 各PCでは Google Maps Clinic Collector v5.7.0 だけを実行してください。clinics.sqlite3 は共有・コピー・書込みしません。\n3. 各PCの結果CSVは上書きせず、PC番号が分かる名前で保存してください。\n4. 3台完了後、結果CSV3つと national_kouseikyoku_master.csv をメインPCに戻します。\n5. 厚生局マスターを先に統合し、その後Maps結果を統合します。\n\nこのキュー作成処理は本番SQLiteを読み取り専用で参照し、既にMaps調査完了済みの医院を除外しています。\n"""
    files["manifest.json"] = json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8")
    files["README_FIRST.txt"] = readme.encode("utf-8-sig")
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in files.items():
            zf.writestr(name, data)
    return out.getvalue(), manifest, files


def collect_all_official(
    cache_dir: Path, progress=None, force: bool = False,
    manual_payloads: dict[str, bytes] | None = None, sources=None,
) -> tuple[pd.DataFrame, list[dict]]:
    """全公式ソースを取得・解析する。

    1地域で止めず、指定された全ソースを最後まで検証してからエラーをまとめて返す。
    これにより「東北を直したら次は関東」のような逐次修正を避ける。
    ただし1ソースでも不完全なら全国キュー自体は作成しない。
    """
    manual_payloads = manual_payloads or {}
    source_list = tuple(sources or OFFICIAL_SOURCES)
    frames: list[pd.DataFrame] = []
    reports: list[dict] = []
    failures: list[str] = []
    total = len(source_list)
    for idx, source in enumerate(source_list, 1):
        if progress:
            progress(idx - 1, total, f"検証中: {source.label}")
        try:
            if source.source_id in manual_payloads:
                payload = manual_payloads[source.source_id]
                cache_path = Path(cache_dir) / f"manual_{source.source_id}.{'xlsx' if source.kind == 'xlsx' else 'zip'}"
                cache_path.parent.mkdir(parents=True, exist_ok=True)
                cache_path.write_bytes(payload)
                cache_hit = False
                mode = "manual"
            else:
                payload, cache_path, cache_hit = download_source(source, cache_dir, force=force)
                mode = "cache" if cache_hit else "download"
            parsed = parse_official_payload(payload, source)
            frames.append(parsed)
            reports.append({
                "source_id": source.source_id, "bureau": source.bureau, "label": source.label,
                "url": source.url, "index_url": source.index_url, "expected_as_of": source.expected_as_of,
                "mode": mode, "cache_file": cache_path.name, "sha256": hashlib.sha256(payload).hexdigest(),
                "rows": int(len(parsed)), "prefectures": sorted(parsed["prefecture"].unique().tolist()),
                "status": "PASS",
            })
            if progress:
                progress(idx, total, f"PASS: {source.label} ({len(parsed):,}件)")
        except Exception as exc:
            message = f"{source.label}: {exc}"
            failures.append(message)
            reports.append({
                "source_id": source.source_id, "bureau": source.bureau, "label": source.label,
                "url": source.url, "index_url": source.index_url, "expected_as_of": source.expected_as_of,
                "mode": "error", "rows": 0, "prefectures": [], "status": "ERROR",
                "error": str(exc),
            })
            if progress:
                progress(idx, total, f"ERROR（続けて全地域を検証）: {source.label}")

    if failures:
        raise ValueError(
            "全国15ソースを最後まで検証しました。エラーは以下です。\n"
            + "\n".join(f"- {msg}" for msg in failures)
        )
    if not frames:
        raise ValueError("全国データを1件も解析できませんでした。")
    combined = pd.concat(frames, ignore_index=True).fillna("")
    combined = combined.drop_duplicates(subset=["clinic_id"], keep="last").reset_index(drop=True)
    covered = set(combined["prefecture"].unique())
    missing = set(PREFECTURES) - covered
    extra = covered - set(PREFECTURES)
    if missing:
        raise ValueError("全国47都道府県が揃っていません: " + ", ".join(sorted(missing)))
    if extra:
        raise ValueError("想定外の都道府県を検出しました: " + ", ".join(sorted(extra)))
    return combined[MASTER_HEADERS], reports


def source_registry_frame() -> pd.DataFrame:
    return pd.DataFrame([
        {
            "地方厚生局": s.bureau, "対象": s.label, "基準日": s.expected_as_of,
            "公式一覧ページ": s.index_url, "自動取得URL": s.url,
        }
        for s in OFFICIAL_SOURCES
    ])


def source_registry_json() -> str:
    return json.dumps([asdict(s) for s in OFFICIAL_SOURCES], ensure_ascii=False, indent=2)
