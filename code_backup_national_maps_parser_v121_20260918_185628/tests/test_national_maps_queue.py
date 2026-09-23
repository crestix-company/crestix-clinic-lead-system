from io import BytesIO
from pathlib import Path
import json
import zipfile

import pandas as pd
from openpyxl import Workbook

from src.master.google_maps import QUEUE_HEADERS
from src.master.national_maps_queue import (
    MASTER_HEADERS, OFFICIAL_SOURCES, PREFECTURES, OfficialSource, build_maps_population, build_package,
    parse_code_workbook, split_three,
)


def make_book(pref="大阪府"):
    book = Workbook(); ws = book.active
    ws.title = pref
    ws.append(["コード内容別医療機関一覧表"])
    ws.append(["[令和 8年 9月 1日現在　医科　現存/休止]"])
    ws.append(["1", "01,1234,5", "見本クリニック", "〒530-0001 大阪市北区梅田1-1", "06-1234-5678\n常勤:1", "開設 太郎", "院長 花子", "令2.1.1\n新規\n令2.1.1", "内 消", "診療所\n現存"])
    ws.append(["2", "01,1234,6", "見本病院", "〒530-0002 大阪市北区梅田1-2", "06-1234-5679", "A", "B", "令2.1.1", "一般 20 内", "病院\n現存"])
    ws.append(["3", "01,1234,7", "見本医療センター", "〒530-0003 大阪市北区梅田1-3", "06-1234-5680", "A", "B", "令2.1.1", "内", "診療所\n現存"])
    ws.append(["4", "01,1234,8", "休止クリニック", "〒530-0004 大阪市北区梅田1-4", "06-1234-5681", "A", "B", "令2.1.1", "内", "診療所\n休止"])
    out = BytesIO(); book.save(out); return out.getvalue()


def test_parse_generic_code_workbook_multiline():
    source = OfficialSource("x", "近畿厚生局", "大阪", "https://kouseikyoku.mhlw.go.jp/x.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="大阪府", expected_prefectures=("大阪府",))
    frame = parse_code_workbook(make_book(), "大阪医科.xlsx", source)
    assert len(frame) == 4
    first = frame.iloc[0]
    assert first["clinic_id"] == "大阪府:0112345"
    assert first["phone"] == "06-1234-5678"
    assert first["facility_type"] == "診療所"
    assert first["status"] == "現存"
    assert first["registration_reason"] == "新規"
    assert first["as_of"] == "2026-09-01"


def test_population_filters_and_completed_ids():
    source = OfficialSource("x", "近畿厚生局", "大阪", "https://kouseikyoku.mhlw.go.jp/x.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="大阪府", expected_prefectures=("大阪府",))
    frame = parse_code_workbook(make_book(), "大阪医科.xlsx", source)
    targets, counts = build_maps_population(frame, {"大阪府:0112345"})
    assert targets == []
    assert counts["already_maps_completed"] == 1
    assert counts["hospital"] == 1
    assert counts["center"] == 1
    assert counts["inactive"] == 1


def test_three_way_split_exact_and_no_duplicates():
    records = [
        {"clinic_id": f"東京都:{i:07d}", "clinic_name": f"医院{i}", "phone": "03-0000-0000", "address": "東京都", "prefecture": "東京都", "medical_type": "医科", "facility_type": "診療所", "status": "現存"}
        for i in range(101)
    ]
    groups = split_three(records)
    assert sorted(len(g) for g in groups) == [33, 34, 34]
    keys = [r["clinic_id"] for g in groups for r in g]
    assert len(keys) == len(set(keys)) == 101
    assert [[r["clinic_id"] for r in g] for g in groups] == [[r["clinic_id"] for r in g] for g in split_three(records)]


def test_package_uses_exact_existing_queue_headers():
    rows = [{
        "clinic_id": f"神奈川県:{i:07d}", "clinic_name": f"医院{i}", "address": "神奈川県横浜市", "phone": "045-000-0000",
        "prefecture": "神奈川県", "medical_type": "医科", "facility_type": "診療所", "owner_name": "", "manager_name": "",
        "designation_date": "", "registration_reason": "", "status": "現存", "as_of": "2026-09-01", "source_url": "",
        "designation_period_start": "", "departments": "内", "source_bureau": "関東信越厚生局", "source_file": "x.xlsx",
    } for i in range(5)]
    frame = pd.DataFrame(rows, columns=MASTER_HEADERS)
    package, manifest, files = build_package(frame, set(), [], {"completed": 0})
    assert manifest["pc_counts"] == {"pc1": 2, "pc2": 2, "pc3": 1}
    with zipfile.ZipFile(BytesIO(package)) as zf:
        assert {"maps_queue_pc1.csv", "maps_queue_pc2.csv", "maps_queue_pc3.csv", "maps_queue_all.csv", "national_kouseikyoku_master.csv", "manifest.json", "README_FIRST.txt"}.issubset(zf.namelist())
        header = zf.read("maps_queue_pc1.csv").decode("utf-8-sig").splitlines()[0].split(",")
        assert header == QUEUE_HEADERS


def test_real_tokyo_workbook_if_present():
    path = Path("data/raw/131コード内容別一覧表（医科）東京r0809.xlsx")
    if not path.exists():
        return
    source = OfficialSource("tokyo", "関東信越厚生局", "東京", "https://kouseikyoku.mhlw.go.jp/x.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", expected_prefectures=("東京都",))
    frame = parse_code_workbook(path.read_bytes(), path.name, source)
    assert len(frame) > 10000
    assert set(frame["prefecture"]) == {"東京都"}
    assert set(frame["medical_type"]) == {"医科"}
    assert frame["clinic_id"].str.startswith("東京都:").all()
    assert frame["facility_type"].isin(["病院", "診療所"]).all()


def test_official_registry_covers_all_47_once():
    covered = [pref for source in OFFICIAL_SOURCES for pref in source.expected_prefectures]
    assert len(covered) == 47
    assert len(set(covered)) == 47
    assert set(covered) == set(PREFECTURES)
    assert all(source.url.startswith("https://kouseikyoku.mhlw.go.jp/") for source in OFFICIAL_SOURCES)


def test_parse_tohoku_hyphen_medical_code_and_primary_number_only():
    book = Workbook(); ws = book.active
    ws.title = "青森県"
    ws.append(["コード内容別医療機関一覧表"])
    ws.append(["[令和 8年 9月 1日現在　医科　現存/休止]"])
    ws.append(["1", "01-1024-3", "青森見本クリニック", "〒030-0000 青森市見本1-1", "017-123-4567", "A", "B", "令2.1.1\n新規\n令2.1.1", "内", "診療所\n現存"])
    ws.append(["2", "01-1216-5\n(01-3216-9)", "青森併設見本クリニック", "〒030-0001 青森市見本1-2", "017-123-4568", "A", "B", "令2.1.1", "内", "診療所\n現存"])
    out = BytesIO(); book.save(out)
    source = OfficialSource("tohoku_test", "東北厚生局", "青森県", "https://kouseikyoku.mhlw.go.jp/x.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="青森県", expected_prefectures=("青森県",))
    frame = parse_code_workbook(out.getvalue(), "青森県医科.xlsx", source)
    assert frame["clinic_id"].tolist() == ["青森県:0110243", "青森県:0112165"]


def test_parse_fullwidth_dash_medical_code():
    book = Workbook(); ws = book.active
    ws.title = "福島県"
    ws.append(["コード内容別医療機関一覧表"])
    ws.append(["[令和 8年 9月 1日現在　医科　現存/休止]"])
    ws.append(["1", "02－1340－6", "福島見本クリニック", "〒960-0000 福島市見本1-1", "024-123-4567", "A", "B", "令2.1.1", "内", "診療所\n現存"])
    out = BytesIO(); book.save(out)
    source = OfficialSource("tohoku_test2", "東北厚生局", "福島県", "https://kouseikyoku.mhlw.go.jp/x.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="福島県", expected_prefectures=("福島県",))
    frame = parse_code_workbook(out.getvalue(), "福島県医科.xlsx", source)
    assert frame.iloc[0]["clinic_id"] == "福島県:0213406"


def test_parse_tohoku_plain_digit_medical_code():
    book = Workbook(); ws = book.active
    ws.title = "宮城県"
    ws.append(["コード内容別医療機関一覧表"])
    ws.append(["[令和 8年 9月 1日現在　医科　現存/休止]"])
    ws.append(["1", "0211239", "宮城見本クリニック", "〒980-0000 仙台市見本1-1", "022-123-4567", "A", "B", "令2.1.1", "内", "診療所\n現存"])
    ws.append(["2", "0210017\n(0231209)", "宮城併設見本クリニック", "〒980-0001 仙台市見本1-2", "022-123-4568", "A", "B", "令2.1.1", "内", "診療所\n現存"])
    out = BytesIO(); book.save(out)
    source = OfficialSource("tohoku_test3", "東北厚生局", "宮城県", "https://kouseikyoku.mhlw.go.jp/x.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="宮城県", expected_prefectures=("宮城県",))
    frame = parse_code_workbook(out.getvalue(), "宮城県医科.xlsx", source)
    assert frame["clinic_id"].tolist() == ["宮城県:0211239", "宮城県:0210017"]


def _one_record_book(pref: str, code: str = "01,1234,5") -> bytes:
    book = Workbook(); ws = book.active
    ws.title = "Sheet1"
    ws.append(["コード内容別医療機関一覧表"])
    ws.append(["[令和 8年 9月 1日現在　医科　現存/休止]"])
    ws.append(["1", code, f"{pref}見本クリニック", "〒100-0001 市区町村見本1-1", "03-1234-5678", "A", "B", "令2.1.1", "内", "診療所\n現存"])
    out = BytesIO(); book.save(out); return out.getvalue()


def test_kanto_all_10_prefectures_from_jis_filename_prefixes():
    from src.master.national_maps_queue import parse_official_payload
    kanto = next(s for s in OFFICIAL_SOURCES if s.source_id == "kanto_shinetsu")
    jis = {"茨城県":"08", "栃木県":"09", "群馬県":"10", "埼玉県":"11", "千葉県":"12", "東京都":"13", "神奈川県":"14", "新潟県":"15", "山梨県":"19", "長野県":"20"}
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for pref, code2 in jis.items():
            # 実ファイルと同様に JIS 2桁 + 帳票区分1 + 日本語名。
            zf.writestr(f"{code2}1コード内容別一覧表（医科）{pref}r0809.xlsx", _one_record_book(pref))
    frame = parse_official_payload(out.getvalue(), kanto)
    assert set(frame["prefecture"]) == set(kanto.expected_prefectures)
    assert len(frame) == 10


def test_prefecture_can_be_recovered_when_japanese_zip_name_is_mojibake():
    from src.master.national_maps_queue import _prefecture_from_identifier, _repair_zip_filename
    original = "111コード内容別一覧表（医科）埼玉r0809.xlsx"
    mojibake = original.encode("cp932").decode("cp437")
    assert "埼玉" not in mojibake
    assert "埼玉" in _repair_zip_filename(mojibake)
    assert _prefecture_from_identifier(mojibake, ("埼玉県",)) == "埼玉県"


def test_official_medical_code_separator_variants_seen_across_regions():
    from src.master.national_maps_queue import _extract_medical_code
    samples = {
        "030,176,2": "0301762",   # 埼玉型
        "01-0198-0": "0101980",   # 千葉/東北型
        "010,006.5": "0100065",   # 神奈川型
        "011,120,4": "0111204",   # 新潟型
        "01,0020,2": "0100202",   # 群馬型
        "01・1206・4": "0112064", # 長野型
        "02－1340－6": "0213406", # 全角ハイフン型
        "0211239": "0211239",     # 区切りなし型
    }
    for raw, expected in samples.items():
        assert _extract_medical_code(raw) == expected, raw


def test_jis_prefix_inference_covers_all_47_without_japanese_names():
    from src.master.national_maps_queue import PREFECTURE_TO_JIS, _prefecture_from_identifier
    for pref in PREFECTURES:
        code = PREFECTURE_TO_JIS[pref]
        assert _prefecture_from_identifier(f"{code}1xxxxxxxx_r0809.xlsx", (pref,)) == pref


def test_collect_all_official_aggregates_source_errors(monkeypatch, tmp_path):
    import src.master.national_maps_queue as nmq
    good = OfficialSource("good", "A局", "良い県", "https://kouseikyoku.mhlw.go.jp/good.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="大阪府", expected_prefectures=("大阪府",))
    bad1 = OfficialSource("bad1", "B局", "失敗1", "https://kouseikyoku.mhlw.go.jp/bad1.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="京都府", expected_prefectures=("京都府",))
    bad2 = OfficialSource("bad2", "C局", "失敗2", "https://kouseikyoku.mhlw.go.jp/bad2.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="兵庫県", expected_prefectures=("兵庫県",))
    payloads = {"good": _one_record_book("大阪府")}
    def fake_download(source, cache_dir, force=False, timeout=45):
        if source.source_id in payloads:
            p = tmp_path / (source.source_id + ".xlsx"); p.write_bytes(payloads[source.source_id])
            return payloads[source.source_id], p, False
        raise RuntimeError("test download failure")
    monkeypatch.setattr(nmq, "download_source", fake_download)
    try:
        nmq.collect_all_official(tmp_path, sources=(good, bad1, bad2))
        assert False, "expected combined error"
    except ValueError as exc:
        msg = str(exc)
        assert "失敗1" in msg and "失敗2" in msg
        assert "全国15ソースを最後まで検証" in msg


def test_mojibake_zip_entry_is_parsed_end_to_end():
    from src.master.national_maps_queue import parse_official_payload
    source = OfficialSource(
        "saitama_zip", "関東信越厚生局", "埼玉県テスト",
        "https://kouseikyoku.mhlw.go.jp/x.zip", "https://kouseikyoku.mhlw.go.jp/", "zip",
        expected_prefectures=("埼玉県",),
    )
    original = "111コード内容別一覧表（医科）埼玉r0809.xlsx"
    mojibake = original.encode("cp932").decode("cp437")
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(mojibake, _one_record_book("埼玉県"))
    frame = parse_official_payload(out.getvalue(), source)
    assert len(frame) == 1
    assert frame.iloc[0]["prefecture"] == "埼玉県"


def _variant_record_book(
    clinic_name: str,
    code: str,
    *,
    address: str = "〒100-0001 市区町村見本1-1",
    owner: str = "開設 太郎",
    manager: str = "院長 花子",
    department: str = "内",
    facility_cell: str = "診療所\n現存",
    extra_tail: str = "",
    shifted_code_col: bool = False,
) -> bytes:
    book = Workbook(); ws = book.active
    ws.title = "Sheet1"
    ws.append(["コード内容別医療機関一覧表"])
    ws.append(["[令和 8年 9月 1日現在　医科　現存/休止]"])
    row = ["1", code, clinic_name, address, "03-1234-5678", owner, manager, "令2.1.1\n新規\n令6.1.1", department, facility_cell]
    if extra_tail:
        row.append(extra_tail)
    if shifted_code_col:
        row.insert(1, "")
    ws.append(row)
    out = BytesIO(); book.save(out); return out.getvalue()


def test_v12_medical_code_variants_from_kinki_and_shikoku():
    from src.master.national_maps_queue import _extract_medical_code
    samples = {
        "01-15202": "0115202",                 # 近畿: 福井/大阪など 2区切り
        "01-01813\n(01-61813)": "0101813",    # 近畿: 併設番号つき
        "011 026.5\n市医26": "0110265",       # 四国: 高知型
        "011 038.0": "0110380",               # 四国: 高知型
        "01-14658": "0114658",                # 四国: 徳島型
        "01-0126-1": "0101261",               # 東海北陸型
    }
    for raw, expected in samples.items():
        assert _extract_medical_code(raw) == expected, raw


def test_v12_kinki_roman_filename_prefecture_aliases():
    from src.master.national_maps_queue import _prefecture_from_identifier
    samples = {
        "2026.9_kikanzentai_fukui_ika.xlsx": "福井県",
        "2026.9_kikanzentai_shiga_ika.xlsx": "滋賀県",
        "2026.9_kikanzentai_kyoto_ika.xlsx": "京都府",
        "2026.9_kikanzentai_osaka_ika.xlsx": "大阪府",
        "2026.9_kikanzentai_hyogo_ika.xlsx": "兵庫県",
        "2026.9_kikanzentai_nara_ika.xlsx": "奈良県",
        "2026.9_kikanzentai_wakayama_ika.xlsx": "和歌山県",
    }
    expected = tuple(samples.values())
    for filename, pref in samples.items():
        assert _prefecture_from_identifier(filename, expected) == pref


def test_v12_kinki_two_segment_code_parses_end_to_end():
    source = OfficialSource(
        "kinki_test", "近畿厚生局", "福井県", "https://kouseikyoku.mhlw.go.jp/x.xlsx",
        "https://kouseikyoku.mhlw.go.jp/", "xlsx", expected_prefectures=("福井県",),
    )
    raw = _variant_record_book(
        "赤井内科呼吸器クリニック", "01-15368", address="〒910-0003 福井市松本2-4-16"
    )
    frame = parse_code_workbook(raw, "2026.9_kikanzentai_fukui_ika.xlsx", source)
    assert len(frame) == 1
    assert frame.iloc[0]["clinic_id"] == "福井県:0115368"
    assert frame.iloc[0]["facility_type"] == "診療所"


def test_v12_shikoku_space_dot_code_parses_end_to_end():
    source = OfficialSource(
        "shikoku_test", "四国厚生支局", "高知県", "https://kouseikyoku.mhlw.go.jp/x.xlsx",
        "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="高知県", expected_prefectures=("高知県",),
    )
    raw = _variant_record_book(
        "三和会国吉病院", "011 026.5\n市医26", address="〒780-0901 高知市上町1-3-4",
        department="一般 69 療養 37 内 外", facility_cell="病院\n現存",
    )
    frame = parse_code_workbook(raw, "01_10 高知 医科 コード内容別医療機関一覧表.xlsx", source)
    assert frame.iloc[0]["clinic_id"] == "高知県:0110265"
    assert frame.iloc[0]["facility_type"] == "病院"


def test_v12_facility_marker_can_live_in_later_tail_cell():
    source = OfficialSource(
        "gunma_test", "関東信越厚生局", "群馬県", "https://kouseikyoku.mhlw.go.jp/x.xlsx",
        "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="群馬県", expected_prefectures=("群馬県",),
    )
    raw = _variant_record_book(
        "前橋赤十字病院", "01,1261,1", address="〒371-0811 前橋市朝倉町389-1",
        department="一般 527 内 外", facility_cell="地域支援", extra_tail="病院\n現存",
    )
    frame = parse_code_workbook(raw, "101コード内容別一覧表（医科）群馬r0809.xlsx", source)
    assert frame.iloc[0]["facility_type"] == "病院"
    assert frame.iloc[0]["status"] == "現存"


def test_v12_missing_facility_marker_uses_name_or_beds_without_dropping_workbook():
    source = OfficialSource(
        "facility_test", "関東信越厚生局", "埼玉県", "https://kouseikyoku.mhlw.go.jp/x.xlsx",
        "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="埼玉県", expected_prefectures=("埼玉県",),
    )
    hospital = parse_code_workbook(
        _variant_record_book("埼玉県立小児医療センター", "01,2222,2", department="一般 300 小", facility_cell="現存"),
        "111コード内容別一覧表（医科）埼玉r0809.xlsx", source,
    )
    # センターは後段で除外されるが、20床以上なら施設区分も病院に寄せる。
    assert hospital.iloc[0]["facility_type"] == "病院"

    clinic = parse_code_workbook(
        _variant_record_book(
            "見本内科", "01,2222,3", owner="医療法人 見本病院 理事長 A",
            department="一般 19 内", facility_cell="現存",
        ),
        "111コード内容別一覧表（医科）埼玉r0809.xlsx", source,
    )
    # 開設者名に「病院」があっても末尾根拠と医院名だけで判定するので誤って病院にしない。
    assert clinic.iloc[0]["facility_type"] == "診療所"


def test_v12_center_is_excluded_even_when_facility_was_inferred():
    source = OfficialSource(
        "center_test", "東海北陸厚生局", "三重県", "https://kouseikyoku.mhlw.go.jp/x.xlsx",
        "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="三重県", expected_prefectures=("三重県",),
    )
    frame = parse_code_workbook(
        _variant_record_book("桑名市総合医療センター", "01-0200-1", department="内", facility_cell="現存"),
        "2609（三重医科）コード内容別医療機関一覧表.xlsx", source,
    )
    targets, counts = build_maps_population(frame, set())
    assert targets == []
    assert counts["center"] == 1


def test_v12_shifted_code_column_is_supported():
    source = OfficialSource(
        "shift_test", "四国厚生支局", "徳島県", "https://kouseikyoku.mhlw.go.jp/x.xlsx",
        "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="徳島県", expected_prefectures=("徳島県",),
    )
    raw = _variant_record_book(
        "石本皮フ科", "01-14658", address="〒770-0815 徳島市助任橋1-19", shifted_code_col=True,
    )
    frame = parse_code_workbook(raw, "01_04_徳島 医科 コード内容別医療機関一覧表.xlsx", source)
    assert frame.iloc[0]["clinic_id"] == "徳島県:0114658"
    assert frame.iloc[0]["clinic_name"] == "石本皮フ科"


def _zip_of_named_books(named_books: dict[str, bytes]) -> bytes:
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, raw in named_books.items():
            zf.writestr(name, raw)
    return out.getvalue()


def test_v12_kinki_all_7_prefectures_parse_with_roman_filenames():
    from src.master.national_maps_queue import parse_official_payload
    source = next(s for s in OFFICIAL_SOURCES if s.source_id == "kinki")
    roman = {
        "福井県": "fukui", "滋賀県": "shiga", "京都府": "kyoto", "大阪府": "osaka",
        "兵庫県": "hyogo", "奈良県": "nara", "和歌山県": "wakayama",
    }
    books = {}
    for i, (pref, alias) in enumerate(roman.items(), 1):
        code = f"01-{15000+i:05d}"
        books[f"2026.9_kikanzentai_{alias}_ika.xlsx"] = _variant_record_book(
            f"{pref}見本クリニック", code, address="〒100-0001 市区町村見本1-1"
        )
    frame = parse_official_payload(_zip_of_named_books(books), source)
    assert set(frame["prefecture"]) == set(source.expected_prefectures)
    assert len(frame) == 7


def test_v12_shikoku_all_4_prefectures_parse_mixed_code_styles():
    from src.master.national_maps_queue import parse_official_payload
    source = next(s for s in OFFICIAL_SOURCES if s.source_id == "shikoku")
    specs = {
        "徳島県": ("01_04_徳島 医科 コード内容別医療機関一覧表.xlsx", "01-14658"),
        "香川県": ("01_01_香川 医科 コード内容別医療機関一覧表.xlsx", "01-12001"),
        "愛媛県": ("01_07 愛媛 医科 コード内容別医療機関一覧表.xlsx", "01-13002"),
        "高知県": ("01_10 高知 医科 コード内容別医療機関一覧表.xlsx", "011 026.5\n市医26"),
    }
    books = {
        filename: _variant_record_book(f"{pref}見本クリニック", code, address="〒100-0001 市区町村見本1-1")
        for pref, (filename, code) in specs.items()
    }
    frame = parse_official_payload(_zip_of_named_books(books), source)
    assert set(frame["prefecture"]) == set(source.expected_prefectures)
    assert len(frame) == 4


def test_v12_tokai_hokuriku_all_6_prefectures_no_whole_file_drop_on_facility_tail_variants():
    from src.master.national_maps_queue import parse_official_payload
    source = next(s for s in OFFICIAL_SOURCES if s.source_id == "tokai_hokuriku")
    prefs = ("富山県", "石川県", "岐阜県", "静岡県", "愛知県", "三重県")
    books = {}
    for i, pref in enumerate(prefs, 1):
        # 地域支援/病院/現存が別セルでもファイル全体を破棄しないことを確認。
        books[f"2609（{pref[:-1]}医科）コード内容別医療機関一覧表.xlsx"] = _variant_record_book(
            f"{pref}赤十字病院", f"01-{120+i:04d}-{i}", address="〒100-0001 市区町村見本1-1",
            department="一般 120 内 外", facility_cell="地域支援", extra_tail="病院\n現存",
        )
    frame = parse_official_payload(_zip_of_named_books(books), source)
    assert set(frame["prefecture"]) == set(source.expected_prefectures)
    assert len(frame) == 6
    assert set(frame["facility_type"]) == {"病院"}
