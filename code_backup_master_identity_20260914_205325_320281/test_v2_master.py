from datetime import date
from io import BytesIO
import json
import sqlite3
import pandas as pd
import pytest
from openpyxl import load_workbook
from src.master.store import ClinicStore
from src.master.filters import Filters
from src.master.comdesk import COMDESK_HEADERS, infer_comdesk_columns
from src.master.samples import load_demo,sample_records
from src.io.input_loader import load_table
from src.io.output_writer import csv_bytes


@pytest.fixture
def store(tmp_path):
    return ClinicStore(tmp_path/"clinics.sqlite3")


def table(records=None,headers=None):
    return load_table(csv_bytes(headers or ["UUID","医院名","電話番号","住所","メモ","メモ", ""],records or [
        ["A","青空内視鏡クリニック","03-0000-0001","東京都千代田区架空町1-1-1"," NA ","=1+1",""]]),"test.csv")


@pytest.mark.parametrize("master_first",[False,True])
def test_uuid_priority_both_import_orders_and_original_preservation(store,master_first):
    r = sample_records()[0]
    if master_first:
        store.import_master([r])
    t = table()
    store.import_comdesk(t)
    if not master_first:
        store.import_master([r])
    assert store.count()==1
    clinic = store.query()[0]
    assert clinic["uuid"]=="A"
    store.save_research(clinic["id"],{"hp_status":"VERIFIED","hp_url":"https://demo.example/"})
    files = store.export(Filters(active_only=False,hp_only=False))
    result = load_table(files["final_comdesk_import.csv"],"final.csv")
    assert result.headers==COMDESK_HEADERS
    assert result.value(0,0)=="A"
    assert result.value(0,2)==t.value(0,1)
    assert result.value(0,9)==t.value(0,2)
    with store.connect() as connection:
        original = connection.execute("SELECT row_json FROM comdesk_original_rows").fetchone()[0]
    assert json.loads(original)==t.data.values.tolist()[0]
    book = load_workbook(BytesIO(files["final_comdesk_import.xlsx"]))
    ws = book.active
    assert ws.max_column==28
    assert ws["A2"].value=="A" and ws["J2"].value=="03-0000-0001"


def test_master_adds_new_and_new_export_uuid_empty(store):
    store.import_comdesk(table())
    result = store.import_master(sample_records())
    assert result=={"MATCHED":1,"NEW":3,"AMBIGUOUS":0}
    assert store.count()==4
    output = load_table(store.export(Filters(hp_only=False))["final_comdesk_import.csv"],"out.csv")
    assert output.value(0,0)=="A"
    assert all(output.value(i,0)=="" for i in [1,2,3])
    assert output.value(1,2)==sample_records()[1]["clinic_name"]
    assert output.value(1,26)==sample_records()[1]["manager_name"]


@pytest.mark.parametrize("master_first", [False, True])
def test_comdesk_28_columns_split_address_uuid_and_new_export(store, master_first):
    records = sample_records()[:2]
    original = ["UUID-00001", "医科", records[0]["clinic_name"], "アオゾラ", "001-0001", "東京都",
                "千代田区架空町1-1-1", "架空ビル 2F", "", records[0]["phone"], "03-9999-0002", "", "",
                "03-9999-0003", "https://demo.example/", "  元の備考\n改行  ", "旧医院名", "既存", "架電済み",
                "記事名", "水", "月火木金土", "09:00", "12:00", "15:00", "18:00", "見本 1郎", "2023/04/01"]
    records[0]["address"] += " 架空ビル 2F"
    incoming = load_table(csv_bytes(COMDESK_HEADERS, [original]), "comdesk.csv")
    inferred = infer_comdesk_columns(incoming)
    assert inferred["clinic_name"] == 2 and inferred["phone"] == 9
    if master_first:
        store.import_master(records)
    store.import_comdesk(incoming)
    if not master_first:
        store.import_master(records)
    assert store.count() == 2 and not store.reviews()
    clinic = next(r for r in store.query() if r["uuid"])
    assert clinic["uuid"] == "UUID-00001" and clinic["address"] == records[0]["address"]
    assert clinic["prefecture"] == "東京都"
    files = store.export(Filters(hp_only=False))
    for filename, raw in files.items():
        exported = load_table(raw, filename)
        assert exported.headers == COMDESK_HEADERS and len(exported.headers) == 28
        assert exported.data.iloc[0].tolist() == original
        new = exported.data.iloc[1].tolist()
        assert new[0] == "" and new[2] == records[1]["clinic_name"]
        assert new[5] == "東京都" and new[6] == records[1]["address"].removeprefix("東京都")
        assert new[7] == "" and new[9] == records[1]["phone"]
        assert all(new[i] == "" for i in (4, 10, 11, 12, 13, 15, 18, 27))


def test_split_address_not_prefixed_twice_and_conflict_is_atomic(store):
    headers = ["UUID", "名前", "Tel1", "都道府県", "住所１", "住所２"]
    incoming = load_table(csv_bytes(headers, [["A", "架空医院", "0312345678", "東京都", "東京都架空区1-1", ""]]), "x.csv")
    store.import_comdesk(incoming)
    assert store.query()[0]["address"] == "東京都架空区1-1"
    invalid = load_table(csv_bytes(headers, [["B", "別の架空医院", "0399990000", "大阪府", "東京都架空区2-2", ""]]), "x.csv")
    with pytest.raises(ValueError, match="都道府県"):
        store.import_comdesk(invalid)
    assert store.count() == 1


def test_exact_name_address_matches_changed_phone(store):
    store.import_comdesk(table())
    record = {**sample_records()[0],"phone":"03-9999-8888"}
    assert store.import_master([record])["MATCHED"]==1
    assert store.query()[0]["uuid"]=="A"
    assert load_table(store.export(Filters(hp_only=False))["final_comdesk_import.csv"],"x.csv").value(0,9)=="03-0000-0001"


def test_ambiguous_shared_phone_held_and_resolved_separate(store):
    records = sample_records()
    records[1]["phone"] = records[0]["phone"]
    result = store.import_master(records[:2])
    assert result["AMBIGUOUS"]==1 and store.count()==1
    review = store.reviews()[0]
    store.resolve_review(review["id"],note="同じ受付番号を共有している別医院")
    assert store.count()==2 and not store.reviews()


def test_fuzzy_not_auto_merged(store):
    r = sample_records()[0]
    store.import_master([r])
    r2 = {**r,"clinic_id":"new-id","clinic_name":"青空内視鏡クリニック東京","phone":"03-1111-2222","address":r["address"].replace("1-1-1","1-1-2")}
    assert store.import_master([r2])["AMBIGUOUS"]==1
    assert store.count()==1


def test_conflicting_uuids_not_lost(store):
    first = table()
    store.import_comdesk(first)
    second = table([["B",*first.data.iloc[0].tolist()[1:]]])
    assert store.import_comdesk(second)["AMBIGUOUS"]==1
    with pytest.raises(ValueError,match="異なる既存UUID"):
        store.resolve_review(store.reviews()[0]["id"],store.query()[0]["id"])
    store.resolve_review(store.reviews()[0]["id"],note="コムデスクで別UUIDであることを確認")
    assert {r["uuid"] for r in store.query()}=={"A","B"}


def test_duplicate_import_idempotent_and_persistent(store):
    assert store.import_comdesk(table())["NEW"]==1
    assert store.import_comdesk(table())["already_imported"]
    store.import_master(sample_records())
    assert store.import_master(sample_records())["already_imported"]
    reopened = ClinicStore(store.path)
    assert reopened.count()==4 and reopened.query()[0]["uuid"]=="A"


def test_monthly_master_diff_and_older_update_rejected(store):
    first = [{**r,"as_of":"2026-08-01"} for r in sample_records()[:2]]
    store.import_master(first)
    second = [{**r,"as_of":"2026-09-01"} for r in sample_records()[:3]]
    result = store.import_master(second)
    assert result["MATCHED"]==2 and result["NEW"]==1
    assert [r["clinic_name"] for r in store.query(Filters(active_only=False,hp_only=False,new_only=True))]==[second[2]["clinic_name"]]
    with pytest.raises(ValueError,match="古い"):
        store.import_master([{**first[0],"phone":"03-0000-9999"}])
    assert store.count()==3


def test_bad_import_atomic_and_no_secret_setting(store):
    bad = table([["A","正常","0312345678","東京都テスト区1-1-1","","",""],["B","","0322223333","東京都テスト区1-1-2","","",""]])
    with pytest.raises(ValueError,match="空白"):
        store.import_comdesk(bad)
    assert store.count()==0
    with pytest.raises(ValueError,match="保存できません"):
        store.set_setting("TAVILY_API_KEY","test-key")


def test_manual_overrides_priority_history_and_reset(store):
    store.import_master(sample_records()[:1]);cid=store.query()[0]["id"]
    store.save_research(cid,{"hp_url":"https://first.example/","hp_status":"VERIFIED","hp_rank":"C"})
    store.override(cid,"hp_rank","A","確認した")
    store.override(cid,"epark_contract","PAID","契約確認URL")
    store.save_research(cid,{"hp_rank":"D","epark_contract":"UNKNOWN"})
    assert store.get(cid)["hp_rank"]=="A"
    assert store.get(cid)["marketing_signal_count"]==1
    assert store.get(cid)["epark_contract"]=="PAID"
    store.override(cid,"hp_rank",None,"自動へ戻す")
    assert store.get(cid)["hp_rank"]=="D"
    assert any(r["action"]=="手動修正:hp_rank" for r in store.history(cid))


def test_director_change_invalidates_previous_age_evidence(store):
    record={**sample_records()[0],"as_of":"2026-08-01"}
    store.import_master([record]);cid=store.query()[0]["id"]
    store.save_research(cid,{"doctor_name":record["manager_name"],"license_registration_year":2005,"age_probability_under_59":.99})
    store.import_master([{**record,"manager_name":"新しい 院長","as_of":"2026-09-01"}])
    assert store.get(cid)["age_probability_under_59"] is None
    assert store.get(cid).get("license_registration_year") is None


def test_manual_rank_remains_even_if_later_search_fails_to_find_hp(store):
    store.import_master(sample_records()[:1]);cid=store.query()[0]["id"]
    store.override(cid,"hp_rank","A","目視確認済み")
    store.save_research(cid,{"hp_status":"NOT_FOUND","hp_rank":"NO_HP"})
    assert store.get(cid)["hp_rank"]=="A"
    assert store.count(Filters())==0


def test_filter_combination_funnel_and_half_boundary(store):
    load_demo(store)
    f = Filters(recent_only=True,age_min=.5,ranks=["A","B"],medical_types=["医科"],prefectures=["東京都"],departments=["消化器内科"],treatments=["内視鏡"],signal_min=2)
    assert store.count(f)==1
    assert store.query(f)[0]["uuid"]=="DEMO-001"
    funnel = store.funnel(f)
    assert funnel[0][1]==4 and funnel[-1][1]==1
    assert all(a[1]>=b[1] for a,b in zip(funnel,funnel[1:]))
    cid = store.query(f)[0]["id"]
    store.save_research(cid,{"age_probability_under_59":.5})
    assert store.count(f)==1
    store.save_research(cid,{"age_probability_under_59":.499999})
    assert store.count(f)==0


@pytest.mark.parametrize("updates,count",[
    ({"signals":["YouTube公式運用","AIチャット導入"]},1),
    ({"treatments":["内視鏡","白内障"]},2),
    ({"uuid_mode":"あり"},1),({"owner_equal":"不一致のみ"},0),({"ranks":["NO_HP"],"hp_only":False},1),
    ({"keyword":"青空"},1),({"keyword":"00000001"},1),({"keyword":"%"},0)])
def test_filter_variants(store,updates,count):
    load_demo(store)
    assert store.count(Filters(**updates))==count


def test_existing_incompatible_database_not_destroyed(tmp_path):
    p = tmp_path/"legacy.db"
    with sqlite3.connect(p) as c:
        c.execute("CREATE TABLE clinics(id INTEGER, name TEXT)")
        c.execute("INSERT INTO clinics VALUES(1,'大切なデータ')")
    with pytest.raises(ValueError,match="形式が異なります"):
        ClinicStore(p)
    with sqlite3.connect(p) as c:
        assert c.execute("SELECT name FROM clinics").fetchone()[0]=="大切なデータ"


def test_backups_and_schema_migration_preserve_originals(store,tmp_path):
    store.import_comdesk(table())
    with store.connect() as c:
        c.execute("PRAGMA user_version=1")
    ClinicStore(store.path)
    assert Path(str(store.path)+".before-v2.bak").exists()
    backup = tmp_path/"restored.db"
    backup.write_bytes(store.backup_bytes())
    assert ClinicStore(backup).query()[0]["uuid"]=="A"


from pathlib import Path
