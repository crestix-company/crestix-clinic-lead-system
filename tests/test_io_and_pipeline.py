from datetime import date
from io import BytesIO
import csv
import io
import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from src.io.input_loader import load_table, infer_columns, sheet_names
from src.io.output_writer import final_csv, build_outputs
from src.enrichment.doctor_license import DoctorLicenseCache
from src.enrichment.epark_checker import EparkChecker
from src.utils.config import ROOT
from src.utils.date_utils import within_years, parse_date
from src.pipeline import run_pipeline


@pytest.mark.parametrize("encoding", ["utf-8-sig","cp932","shift_jis"])
def test_roundtrip_preserves_duplicate_headers_cells_and_order(encoding):
    raw = '電話番号,医院名,住所,その他,その他,\r\n03-0000-0001,架空医院,"東京都,架空区"," 改行\nあり ",NA,\r\n03-0000-0002,別医院,,,NULL,\r\n'.encode(encoding)
    table = load_table(raw,"leads.csv")
    output = final_csv(table,[0],encoding)
    before = list(csv.reader(io.StringIO(raw.decode(encoding))))
    after = list(csv.reader(io.StringIO(output.decode(encoding))))
    assert after == before[:2]
    assert table.headers[-1] == ""
    assert len(table.headers) == 6


def test_xlsx_zero_format_multiple_sheets_and_formulas():
    book = Workbook(); ws = book.active; ws.title="営業"
    ws.append(["電話番号","医院名"]); ws.append([300000001,"架空医院"])
    ws["A2"].number_format = "0000000000"
    ws["D1"].number_format = "@"  # 値のない末尾列もテンプレートの列として保持
    book.create_sheet("補足")
    raw = BytesIO(); book.save(raw)
    assert sheet_names(raw.getvalue(),"x.xlsx") == ["営業","補足"]
    table = load_table(raw.getvalue(),"x.xlsx",sheet="営業")
    assert table.value(0,0) == "0300000001"
    assert table.headers == ["電話番号", "医院名", "", ""]
    assert len(table.data.columns) == 4
    ws["A2"]="=1+2"; raw=BytesIO();book.save(raw)
    with pytest.raises(ValueError,match="数式"):
        load_table(raw.getvalue(),"x.xlsx")


def test_bad_csv_and_empty_file_are_clear_errors():
    with pytest.raises(ValueError,match="列数"):
        load_table(b"A,B\n1,2,3\n","x.csv")
    with pytest.raises(ValueError,match="ヘッダー"):
        load_table(b"","x.csv")
    with pytest.raises(ValueError,match="文字コード"):
        load_table(b"\xff\xfe\x00","x.csv")


def test_date_boundaries_and_era():
    assert within_years("2016-09-10",date(2026,9,10)) is True
    assert within_years("2016-09-09",date(2026,9,10)) is False
    assert within_years("2027-01-01",date(2026,9,10)) is None
    assert within_years("",date(2026,9,10)) is None
    assert within_years("2016-02-29",date(2026,2,28)) is True
    assert parse_date("令8. 9. 1") == date(2026,9,1)
    assert parse_date("昭32. 11. 1") == date(1957,11,1)


def sample_result(config):
    path=ROOT/"samples"
    read=lambda name: pd.read_csv(path/name,dtype=str,keep_default_na=False)
    table=load_table((path/"sample_comdesk.csv").read_bytes(),"sample.csv")
    result=run_pipeline(table,infer_columns(table),read("sample_kouseikyoku.csv"),config,
                        DoctorLicenseCache(frame=read("sample_doctor_license.csv")),
                        EparkChecker(config["epark"],frame=read("sample_epark_confirmed.csv"),as_of=date(2026,9,10)),
                        as_of=date(2026,9,10),extra=read("sample_enrichment.csv"))
    return table,result


def test_end_to_end_expected_leads_and_exact_export(config):
    table, result=sample_result(config)
    assert result.final_indices == [0,1,2,3,12,13]
    assert result.excluded_indices == [4,6,7,8]
    assert 10 in result.review_indices and 12 in result.review_indices
    files=build_outputs(table,result)
    assert set(files)=={"final_comdesk_import.csv","judged_all.xlsx","excluded.csv","needs_review.csv"}
    restored=load_table(files["final_comdesk_import.csv"],"final.csv")
    assert restored.headers==table.headers
    assert restored.data.values.tolist()==table.data.iloc[result.final_indices].values.tolist()
    book=load_workbook(BytesIO(files["judged_all.xlsx"]),data_only=True)
    ws=book["判定付き全件"]
    assert ws.max_row==17
    headers=[c.value for c in ws[1]]
    assert isinstance(ws.cell(2,headers.index("59歳以下確率")+1).value,float)
    assert ws.cell(2,1).value=="03-0000-0001"


def test_no_master_is_review_not_dropped(config):
    table=load_table("電話番号,医院名\n03-0000-0001,架空医院\n".encode(),"x.csv")
    result=run_pipeline(table,infer_columns(table),pd.DataFrame(),config,
                        DoctorLicenseCache(),EparkChecker(config["epark"]),as_of=date(2026,9,10))
    assert result.final_indices==[]
    assert result.excluded_indices==[]
    assert result.review_indices==[0]
