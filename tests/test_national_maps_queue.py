from io import BytesIO
from pathlib import Path
import json
import zipfile

import pandas as pd
from openpyxl import Workbook, load_workbook

from src.master.google_maps import QUEUE_HEADERS
from src.master.national_maps_queue import (
    MASTER_HEADERS, OFFICIAL_SOURCES, PREFECTURES, OfficialSource, build_maps_population, build_package,
    parse_code_workbook, split_three,
)


def make_book(pref="大阪府"):
    book = Workbook(); ws = book.active
    ws.title = pref
    ws.append(["コード内容別医療機関一覧表"])
    ws.append(["[令和 8年 9月 1日現在　医科　現存/休止]"])
    ws.append(["1", "01,1234,5", "見本クリニック", "〒530-0001 大阪市北区梅田1-1", "06-1234-5678\n常勤:1", "開設 太郎", "院長 花子", "令2.1.1\n新規\n令2.1.1", "内 消", "診療所\n現存"])
    ws.append(["2", "01,1234,6", "見本病院", "〒530-0002 大阪市北区梅田1-2", "06-1234-5679", "A", "B", "令2.1.1", "一般 20 内", "病院\n現存"])
    ws.append(["3", "01,1234,7", "見本医療センター", "〒530-0003 大阪市北区梅田1-3", "06-1234-5680", "A", "B", "令2.1.1", "内", "診療所\n現存"])
    ws.append(["4", "01,1234,8", "休止クリニック", "〒530-0004 大阪市北区梅田1-4", "06-1234-5681", "A", "B", "令2.1.1", "内", "診療所\n休止"])
    out = BytesIO(); book.save(out); return out.getvalue()


def test_parse_generic_code_workbook_multiline():
    source = OfficialSource("x", "近畿厚生局", "大阪", "https://kouseikyoku.mhlw.go.jp/x.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="大阪府", expected_prefectures=("大阪府",))
    frame = parse_code_workbook(make_book(), "大阪医科.xlsx", source)
    assert len(frame) == 4
    first = frame.iloc[0]
    assert first["clinic_id"] == "大阪府:0112345"
    assert first["phone"] == "06-1234-5678"
    assert first["facility_type"] == "診療所"
    assert first["status"] == "現存"
    assert first["registration_reason"] == "新規"
    assert first["as_of"] == "2026-09-01"


def test_population_filters_and_completed_ids():
    source = OfficialSource("x", "近畿厚生局", "大阪", "https://kouseikyoku.mhlw.go.jp/x.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="大阪府", expected_prefectures=("大阪府",))
    frame = parse_code_workbook(make_book(), "大阪医科.xlsx", source)
    targets, counts = build_maps_population(frame, {"大阪府:0112345"})
    assert targets == []
    assert counts["already_maps_completed"] == 1
    assert counts["hospital"] == 1
    assert counts["center"] == 1
    assert counts["inactive"] == 1


def test_three_way_split_exact_and_no_duplicates():
    records = [
        {"clinic_id": f"東京都:{i:07d}", "clinic_name": f"医院{i}", "phone": "03-0000-0000", "address": "東京都", "prefecture": "東京都", "medical_type": "医科", "facility_type": "診療所", "status": "現存"}
        for i in range(101)
    ]
    groups = split_three(records)
    assert sorted(len(g) for g in groups) == [33, 34, 34]
    keys = [r["clinic_id"] for g in groups for r in g]
    assert len(keys) == len(set(keys)) == 101
    assert [[r["clinic_id"] for r in g] for g in groups] == [[r["clinic_id"] for r in g] for g in split_three(records)]


def test_package_uses_exact_existing_queue_headers():
    rows = [{
        "clinic_id": f"神奈川県:{i:07d}", "clinic_name": f"医院{i}", "address": "神奈川県横浜市", "phone": "045-000-0000",
        "prefecture": "神奈川県", "medical_type": "医科", "facility_type": "診療所", "owner_name": "", "manager_name": "",
        "designation_date": "", "registration_reason": "", "status": "現存", "as_of": "2026-09-01", "source_url": "",
        "designation_period_start": "", "departments": "内", "source_bureau": "関東信越厚生局", "source_file": "x.xlsx",
    } for i in range(5)]
    frame = pd.DataFrame(rows, columns=MASTER_HEADERS)
    package, manifest, files = build_package(frame, set(), [], {"completed": 0})
    assert manifest["pc_counts"] == {"pc1": 2, "pc2": 2, "pc3": 1}
    with zipfile.ZipFile(BytesIO(package)) as zf:
        assert {"maps_queue_pc1.csv", "maps_queue_pc2.csv", "maps_queue_pc3.csv", "maps_queue_all.csv", "national_kouseikyoku_master.csv", "manifest.json", "README_FIRST.txt"}.issubset(zf.namelist())
        header = zf.read("maps_queue_pc1.csv").decode("utf-8-sig").splitlines()[0].split(",")
        assert header == QUEUE_HEADERS


def test_real_tokyo_workbook_if_present():
    path = Path("data/raw/131コード内容別一覧表（医科）東京r0809.xlsx")
    if not path.exists():
        return
    source = OfficialSource("tokyo", "関東信越厚生局", "東京", "https://kouseikyoku.mhlw.go.jp/x.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", expected_prefectures=("東京都",))
    frame = parse_code_workbook(path.read_bytes(), path.name, source)
    assert len(frame) > 10000
    assert set(frame["prefecture"]) == {"東京都"}
    assert set(frame["medical_type"]) == {"医科"}
    assert frame["clinic_id"].str.startswith("東京都:").all()
    assert frame["facility_type"].isin(["病院", "診療所"]).all()


def test_official_registry_covers_all_47_once_per_medical_type():
    from src.master.national_maps_queue import sources_for_medical_types
    assert len(OFFICIAL_SOURCES) == 30
    for medical_type in ("医科", "歯科"):
        sources = sources_for_medical_types([medical_type])
        assert len(sources) == 15
        covered = [pref for source in sources for pref in source.expected_prefectures]
        assert len(covered) == 47
        assert len(set(covered)) == 47
        assert set(covered) == set(PREFECTURES)
        assert all(source.medical_type_hint == medical_type for source in sources)
    assert all(source.url.startswith("https://kouseikyoku.mhlw.go.jp/") for source in OFFICIAL_SOURCES)


def test_parse_tohoku_hyphen_medical_code_and_primary_number_only():
    book = Workbook(); ws = book.active
    ws.title = "青森県"
    ws.append(["コード内容別医療機関一覧表"])
    ws.append(["[令和 8年 9月 1日現在　医科　現存/休止]"])
    ws.append(["1", "01-1024-3", "青森見本クリニック", "〒030-0000 青森市見本1-1", "017-123-4567", "A", "B", "令2.1.1\n新規\n令2.1.1", "内", "診療所\n現存"])
    ws.append(["2", "01-1216-5\n(01-3216-9)", "青森併設見本クリニック", "〒030-0001 青森市見本1-2", "017-123-4568", "A", "B", "令2.1.1", "内", "診療所\n現存"])
    out = BytesIO(); book.save(out)
    source = OfficialSource("tohoku_test", "東北厚生局", "青森県", "https://kouseikyoku.mhlw.go.jp/x.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="青森県", expected_prefectures=("青森県",))
    frame = parse_code_workbook(out.getvalue(), "青森県医科.xlsx", source)
    assert frame["clinic_id"].tolist() == ["青森県:0110243", "青森県:0112165"]


def test_parse_fullwidth_dash_medical_code():
    book = Workbook(); ws = book.active
    ws.title = "福島県"
    ws.append(["コード内容別医療機関一覧表"])
    ws.append(["[令和 8年 9月 1日現在　医科　現存/休止]"])
    ws.append(["1", "02－1340－6", "福島見本クリニック", "〒960-0000 福島市見本1-1", "024-123-4567", "A", "B", "令2.1.1", "内", "診療所\n現存"])
    out = BytesIO(); book.save(out)
    source = OfficialSource("tohoku_test2", "東北厚生局", "福島県", "https://kouseikyoku.mhlw.go.jp/x.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="福島県", expected_prefectures=("福島県",))
    frame = parse_code_workbook(out.getvalue(), "福島県医科.xlsx", source)
    assert frame.iloc[0]["clinic_id"] == "福島県:0213406"


def test_parse_tohoku_plain_digit_medical_code():
    book = Workbook(); ws = book.active
    ws.title = "宮城県"
    ws.append(["コード内容別医療機関一覧表"])
    ws.append(["[令和 8年 9月 1日現在　医科　現存/休止]"])
    ws.append(["1", "0211239", "宮城見本クリニック", "〒980-0000 仙台市見本1-1", "022-123-4567", "A", "B", "令2.1.1", "内", "診療所\n現存"])
    ws.append(["2", "0210017\n(0231209)", "宮城併設見本クリニック", "〒980-0001 仙台市見本1-2", "022-123-4568", "A", "B", "令2.1.1", "内", "診療所\n現存"])
    out = BytesIO(); book.save(out)
    source = OfficialSource("tohoku_test3", "東北厚生局", "宮城県", "https://kouseikyoku.mhlw.go.jp/x.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="宮城県", expected_prefectures=("宮城県",))
    frame = parse_code_workbook(out.getvalue(), "宮城県医科.xlsx", source)
    assert frame["clinic_id"].tolist() == ["宮城県:0211239", "宮城県:0210017"]


def _one_record_book(pref: str, code: str = "01,1234,5") -> bytes:
    book = Workbook(); ws = book.active
    ws.title = "Sheet1"
    ws.append(["コード内容別医療機関一覧表"])
    ws.append(["[令和 8年 9月 1日現在　医科　現存/休止]"])
    ws.append(["1", code, f"{pref}見本クリニック", "〒100-0001 市区町村見本1-1", "03-1234-5678", "A", "B", "令2.1.1", "内", "診療所\n現存"])
    out = BytesIO(); book.save(out); return out.getvalue()


def test_kanto_all_10_prefectures_from_jis_filename_prefixes():
    from src.master.national_maps_queue import parse_official_payload
    kanto = next(s for s in OFFICIAL_SOURCES if s.source_id == "kanto_shinetsu")
    jis = {"茨城県":"08", "栃木県":"09", "群馬県":"10", "埼玉県":"11", "千葉県":"12", "東京都":"13", "神奈川県":"14", "新潟県":"15", "山梨県":"19", "長野県":"20"}
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for pref, code2 in jis.items():
            # 実ファイルと同様に JIS 2桁 + 帳票区分1 + 日本語名。
            zf.writestr(f"{code2}1コード内容別一覧表（医科）{pref}r0809.xlsx", _one_record_book(pref))
    frame = parse_official_payload(out.getvalue(), kanto)
    assert set(frame["prefecture"]) == set(kanto.expected_prefectures)
    assert len(frame) == 10


def test_prefecture_can_be_recovered_when_japanese_zip_name_is_mojibake():
    from src.master.national_maps_queue import _prefecture_from_identifier, _repair_zip_filename
    original = "111コード内容別一覧表（医科）埼玉r0809.xlsx"
    mojibake = original.encode("cp932").decode("cp437")
    assert "埼玉" not in mojibake
    assert "埼玉" in _repair_zip_filename(mojibake)
    assert _prefecture_from_identifier(mojibake, ("埼玉県",)) == "埼玉県"


def test_official_medical_code_separator_variants_seen_across_regions():
    from src.master.national_maps_queue import _extract_medical_code
    samples = {
        "030,176,2": "0301762",   # 埼玉型
        "01-0198-0": "0101980",   # 千葉/東北型
        "010,006.5": "0100065",   # 神奈川型
        "011,120,4": "0111204",   # 新潟型
        "01,0020,2": "0100202",   # 群馬型
        "01・1206・4": "0112064", # 長野型
        "02－1340－6": "0213406", # 全角ハイフン型
        "0211239": "0211239",     # 区切りなし型
    }
    for raw, expected in samples.items():
        assert _extract_medical_code(raw) == expected, raw


def test_jis_prefix_inference_covers_all_47_without_japanese_names():
    from src.master.national_maps_queue import PREFECTURE_TO_JIS, _prefecture_from_identifier
    for pref in PREFECTURES:
        code = PREFECTURE_TO_JIS[pref]
        assert _prefecture_from_identifier(f"{code}1xxxxxxxx_r0809.xlsx", (pref,)) == pref


def test_collect_all_official_aggregates_source_errors(monkeypatch, tmp_path):
    import src.master.national_maps_queue as nmq
    good = OfficialSource("good", "A局", "良い県", "https://kouseikyoku.mhlw.go.jp/good.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="大阪府", expected_prefectures=("大阪府",))
    bad1 = OfficialSource("bad1", "B局", "失敗1", "https://kouseikyoku.mhlw.go.jp/bad1.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="京都府", expected_prefectures=("京都府",))
    bad2 = OfficialSource("bad2", "C局", "失敗2", "https://kouseikyoku.mhlw.go.jp/bad2.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="兵庫県", expected_prefectures=("兵庫県",))
    payloads = {"good": _one_record_book("大阪府")}
    def fake_download(source, cache_dir, force=False, timeout=45):
        if source.source_id in payloads:
            p = tmp_path / (source.source_id + ".xlsx"); p.write_bytes(payloads[source.source_id])
            return payloads[source.source_id], p, False
        raise RuntimeError("test download failure")
    monkeypatch.setattr(nmq, "download_source", fake_download)
    try:
        nmq.collect_all_official(tmp_path, sources=(good, bad1, bad2))
        assert False, "expected combined error"
    except ValueError as exc:
        msg = str(exc)
        assert "失敗1" in msg and "失敗2" in msg
        assert "全国3ソースを最後まで検証" in msg


def test_mojibake_zip_entry_is_parsed_end_to_end():
    from src.master.national_maps_queue import parse_official_payload
    source = OfficialSource(
        "saitama_zip", "関東信越厚生局", "埼玉県テスト",
        "https://kouseikyoku.mhlw.go.jp/x.zip", "https://kouseikyoku.mhlw.go.jp/", "zip",
        expected_prefectures=("埼玉県",),
    )
    original = "111コード内容別一覧表（医科）埼玉r0809.xlsx"
    mojibake = original.encode("cp932").decode("cp437")
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(mojibake, _one_record_book("埼玉県"))
    frame = parse_official_payload(out.getvalue(), source)
    assert len(frame) == 1
    assert frame.iloc[0]["prefecture"] == "埼玉県"


def _variant_record_book(
    clinic_name: str,
    code: str,
    *,
    address: str = "〒100-0001 市区町村見本1-1",
    owner: str = "開設 太郎",
    manager: str = "院長 花子",
    department: str = "内",
    facility_cell: str = "診療所\n現存",
    extra_tail: str = "",
    shifted_code_col: bool = False,
) -> bytes:
    book = Workbook(); ws = book.active
    ws.title = "Sheet1"
    ws.append(["コード内容別医療機関一覧表"])
    ws.append(["[令和 8年 9月 1日現在　医科　現存/休止]"])
    row = ["1", code, clinic_name, address, "03-1234-5678", owner, manager, "令2.1.1\n新規\n令6.1.1", department, facility_cell]
    if extra_tail:
        row.append(extra_tail)
    if shifted_code_col:
        row.insert(1, "")
    ws.append(row)
    out = BytesIO(); book.save(out); return out.getvalue()


def test_v12_medical_code_variants_from_kinki_and_shikoku():
    from src.master.national_maps_queue import _extract_medical_code
    samples = {
        "01-15202": "0115202",                 # 近畿: 福井/大阪など 2区切り
        "01-01813\n(01-61813)": "0101813",    # 近畿: 併設番号つき
        "011 026.5\n市医26": "0110265",       # 四国: 高知型
        "011 038.0": "0110380",               # 四国: 高知型
        "01-14658": "0114658",                # 四国: 徳島型
        "01-0126-1": "0101261",               # 東海北陸型
    }
    for raw, expected in samples.items():
        assert _extract_medical_code(raw) == expected, raw


def test_v12_kinki_roman_filename_prefecture_aliases():
    from src.master.national_maps_queue import _prefecture_from_identifier
    samples = {
        "2026.9_kikanzentai_fukui_ika.xlsx": "福井県",
        "2026.9_kikanzentai_shiga_ika.xlsx": "滋賀県",
        "2026.9_kikanzentai_kyoto_ika.xlsx": "京都府",
        "2026.9_kikanzentai_osaka_ika.xlsx": "大阪府",
        "2026.9_kikanzentai_hyogo_ika.xlsx": "兵庫県",
        "2026.9_kikanzentai_nara_ika.xlsx": "奈良県",
        "2026.9_kikanzentai_wakayama_ika.xlsx": "和歌山県",
    }
    expected = tuple(samples.values())
    for filename, pref in samples.items():
        assert _prefecture_from_identifier(filename, expected) == pref


def test_v12_kinki_two_segment_code_parses_end_to_end():
    source = OfficialSource(
        "kinki_test", "近畿厚生局", "福井県", "https://kouseikyoku.mhlw.go.jp/x.xlsx",
        "https://kouseikyoku.mhlw.go.jp/", "xlsx", expected_prefectures=("福井県",),
    )
    raw = _variant_record_book(
        "赤井内科呼吸器クリニック", "01-15368", address="〒910-0003 福井市松本2-4-16"
    )
    frame = parse_code_workbook(raw, "2026.9_kikanzentai_fukui_ika.xlsx", source)
    assert len(frame) == 1
    assert frame.iloc[0]["clinic_id"] == "福井県:0115368"
    assert frame.iloc[0]["facility_type"] == "診療所"


def test_v12_shikoku_space_dot_code_parses_end_to_end():
    source = OfficialSource(
        "shikoku_test", "四国厚生支局", "高知県", "https://kouseikyoku.mhlw.go.jp/x.xlsx",
        "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="高知県", expected_prefectures=("高知県",),
    )
    raw = _variant_record_book(
        "三和会国吉病院", "011 026.5\n市医26", address="〒780-0901 高知市上町1-3-4",
        department="一般 69 療養 37 内 外", facility_cell="病院\n現存",
    )
    frame = parse_code_workbook(raw, "01_10 高知 医科 コード内容別医療機関一覧表.xlsx", source)
    assert frame.iloc[0]["clinic_id"] == "高知県:0110265"
    assert frame.iloc[0]["facility_type"] == "病院"


def test_v12_facility_marker_can_live_in_later_tail_cell():
    source = OfficialSource(
        "gunma_test", "関東信越厚生局", "群馬県", "https://kouseikyoku.mhlw.go.jp/x.xlsx",
        "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="群馬県", expected_prefectures=("群馬県",),
    )
    raw = _variant_record_book(
        "前橋赤十字病院", "01,1261,1", address="〒371-0811 前橋市朝倉町389-1",
        department="一般 527 内 外", facility_cell="地域支援", extra_tail="病院\n現存",
    )
    frame = parse_code_workbook(raw, "101コード内容別一覧表（医科）群馬r0809.xlsx", source)
    assert frame.iloc[0]["facility_type"] == "病院"
    assert frame.iloc[0]["status"] == "現存"


def test_v12_missing_facility_marker_uses_name_or_beds_without_dropping_workbook():
    source = OfficialSource(
        "facility_test", "関東信越厚生局", "埼玉県", "https://kouseikyoku.mhlw.go.jp/x.xlsx",
        "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="埼玉県", expected_prefectures=("埼玉県",),
    )
    hospital = parse_code_workbook(
        _variant_record_book("埼玉県立小児医療センター", "01,2222,2", department="一般 300 小", facility_cell="現存"),
        "111コード内容別一覧表（医科）埼玉r0809.xlsx", source,
    )
    # センターは後段で除外されるが、20床以上なら施設区分も病院に寄せる。
    assert hospital.iloc[0]["facility_type"] == "病院"

    clinic = parse_code_workbook(
        _variant_record_book(
            "見本内科", "01,2222,3", owner="医療法人 見本病院 理事長 A",
            department="一般 19 内", facility_cell="現存",
        ),
        "111コード内容別一覧表（医科）埼玉r0809.xlsx", source,
    )
    # 開設者名に「病院」があっても末尾根拠と医院名だけで判定するので誤って病院にしない。
    assert clinic.iloc[0]["facility_type"] == "診療所"


def test_v12_center_is_excluded_even_when_facility_was_inferred():
    source = OfficialSource(
        "center_test", "東海北陸厚生局", "三重県", "https://kouseikyoku.mhlw.go.jp/x.xlsx",
        "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="三重県", expected_prefectures=("三重県",),
    )
    frame = parse_code_workbook(
        _variant_record_book("桑名市総合医療センター", "01-0200-1", department="内", facility_cell="現存"),
        "2609（三重医科）コード内容別医療機関一覧表.xlsx", source,
    )
    targets, counts = build_maps_population(frame, set())
    assert targets == []
    assert counts["center"] == 1


def test_v12_shifted_code_column_is_supported():
    source = OfficialSource(
        "shift_test", "四国厚生支局", "徳島県", "https://kouseikyoku.mhlw.go.jp/x.xlsx",
        "https://kouseikyoku.mhlw.go.jp/", "xlsx", prefecture_hint="徳島県", expected_prefectures=("徳島県",),
    )
    raw = _variant_record_book(
        "石本皮フ科", "01-14658", address="〒770-0815 徳島市助任橋1-19", shifted_code_col=True,
    )
    frame = parse_code_workbook(raw, "01_04_徳島 医科 コード内容別医療機関一覧表.xlsx", source)
    assert frame.iloc[0]["clinic_id"] == "徳島県:0114658"
    assert frame.iloc[0]["clinic_name"] == "石本皮フ科"


def _zip_of_named_books(named_books: dict[str, bytes]) -> bytes:
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, raw in named_books.items():
            zf.writestr(name, raw)
    return out.getvalue()


def test_v12_kinki_all_7_prefectures_parse_with_roman_filenames():
    from src.master.national_maps_queue import parse_official_payload
    source = next(s for s in OFFICIAL_SOURCES if s.source_id == "kinki")
    roman = {
        "福井県": "fukui", "滋賀県": "shiga", "京都府": "kyoto", "大阪府": "osaka",
        "兵庫県": "hyogo", "奈良県": "nara", "和歌山県": "wakayama",
    }
    books = {}
    for i, (pref, alias) in enumerate(roman.items(), 1):
        code = f"01-{15000+i:05d}"
        books[f"2026.9_kikanzentai_{alias}_ika.xlsx"] = _variant_record_book(
            f"{pref}見本クリニック", code, address="〒100-0001 市区町村見本1-1"
        )
    frame = parse_official_payload(_zip_of_named_books(books), source)
    assert set(frame["prefecture"]) == set(source.expected_prefectures)
    assert len(frame) == 7


def test_v12_shikoku_all_4_prefectures_parse_mixed_code_styles():
    from src.master.national_maps_queue import parse_official_payload
    source = next(s for s in OFFICIAL_SOURCES if s.source_id == "shikoku")
    specs = {
        "徳島県": ("01_04_徳島 医科 コード内容別医療機関一覧表.xlsx", "01-14658"),
        "香川県": ("01_01_香川 医科 コード内容別医療機関一覧表.xlsx", "01-12001"),
        "愛媛県": ("01_07 愛媛 医科 コード内容別医療機関一覧表.xlsx", "01-13002"),
        "高知県": ("01_10 高知 医科 コード内容別医療機関一覧表.xlsx", "011 026.5\n市医26"),
    }
    books = {
        filename: _variant_record_book(f"{pref}見本クリニック", code, address="〒100-0001 市区町村見本1-1")
        for pref, (filename, code) in specs.items()
    }
    frame = parse_official_payload(_zip_of_named_books(books), source)
    assert set(frame["prefecture"]) == set(source.expected_prefectures)
    assert len(frame) == 4


def test_v12_tokai_hokuriku_all_6_prefectures_no_whole_file_drop_on_facility_tail_variants():
    from src.master.national_maps_queue import parse_official_payload
    source = next(s for s in OFFICIAL_SOURCES if s.source_id == "tokai_hokuriku")
    prefs = ("富山県", "石川県", "岐阜県", "静岡県", "愛知県", "三重県")
    books = {}
    for i, pref in enumerate(prefs, 1):
        # 地域支援/病院/現存が別セルでもファイル全体を破棄しないことを確認。
        books[f"2609（{pref[:-1]}医科）コード内容別医療機関一覧表.xlsx"] = _variant_record_book(
            f"{pref}赤十字病院", f"01-{120+i:04d}-{i}", address="〒100-0001 市区町村見本1-1",
            department="一般 120 内 外", facility_cell="地域支援", extra_tail="病院\n現存",
        )
    frame = parse_official_payload(_zip_of_named_books(books), source)
    assert set(frame["prefecture"]) == set(source.expected_prefectures)
    assert len(frame) == 6
    assert set(frame["facility_type"]) == {"病院"}


def test_v121_prefecture_hint_is_authoritative_inside_prefecture_workbook():
    """県別Workbook内に他県住所/注記があっても、その県へ切り替えない。"""
    source = OfficialSource(
        "kinki_osaka_test", "近畿厚生局", "近畿テスト",
        "https://kouseikyoku.mhlw.go.jp/x.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx",
        prefecture_hint="大阪府",
        expected_prefectures=("福井県", "滋賀県", "京都府", "大阪府", "兵庫県", "奈良県", "和歌山県"),
    )
    book = Workbook(); ws = book.active
    ws.title = "Sheet1"
    ws.append(["コード内容別医療機関一覧表"])
    ws.append(["[令和 8年 9月 1日現在　医科　現存/休止]"])
    # 1件目は通常の大阪施設。
    ws.append(["1", "01-15202", "大阪見本クリニック", "〒530-0001 大阪市北区梅田1-1", "06-1234-5678", "A", "B", "令2.1.1", "内", "診療所\n現存"])
    # 実帳票の注記・開設者本部住所等で他県名が単独行に現れても県切替しない。
    ws.append(["東京都"])
    # 次施設の住所欄に都道府県名がなくても、ファイル由来の大阪府を維持する。
    ws.append(["2", "01-15203", "大阪第二クリニック", "〒530-0002 大阪市北区梅田1-2", "06-1234-5679", "A", "B", "令2.1.1", "内", "診療所\n現存"])
    out = BytesIO(); book.save(out)

    frame = parse_code_workbook(out.getvalue(), "2026.9_kikanzentai_osaka_ika.xlsx", source)
    assert set(frame["prefecture"]) == {"大阪府"}
    assert frame["clinic_id"].str.startswith("大阪府:").all()


def test_v121_kinki_zip_cannot_emit_prefecture_outside_expected_set():
    """近畿ZIPの県別ファイルから東京都など想定外県が混入しない回帰テスト。"""
    from src.master.national_maps_queue import parse_official_payload
    source = next(s for s in OFFICIAL_SOURCES if s.source_id == "kinki")
    roman = {
        "福井県": "fukui", "滋賀県": "shiga", "京都府": "kyoto", "大阪府": "osaka",
        "兵庫県": "hyogo", "奈良県": "nara", "和歌山県": "wakayama",
    }
    books = {}
    for i, (pref, alias) in enumerate(roman.items(), 1):
        raw = _variant_record_book(
            f"{pref}見本クリニック", f"01-{16000+i:05d}",
            address="〒100-0001 市区町村見本1-1",
        )
        # 大阪Workbookだけ、施設データの後ろに他県名の注記行を混ぜる。
        if pref == "大阪府":
            wb = load_workbook(BytesIO(raw))
            ws = wb.active
            ws.append(["東京都"])
            ws.append(["2", "01-16099", "大阪追加クリニック", "〒530-0002 大阪市北区見本2-2", "06-1111-2222", "A", "B", "令2.1.1", "内", "診療所\n現存"])
            buf = BytesIO(); wb.save(buf); wb.close(); raw = buf.getvalue()
        books[f"2026.9_kikanzentai_{alias}_ika.xlsx"] = raw

    frame = parse_official_payload(_zip_of_named_books(books), source)
    assert set(frame["prefecture"]) == set(source.expected_prefectures)
    assert "東京都" not in set(frame["prefecture"])


def _dental_record_book(pref: str = "大阪府", code: str = "01-12345", name: str = "見本歯科クリニック") -> bytes:
    book = Workbook(); ws = book.active
    ws.title = pref
    ws.append(["コード内容別医療機関一覧表"])
    ws.append(["[令和 8年 9月 1日現在　歯科　現存/休止]"])
    ws.append(["1", code, name, f"〒530-0001 {pref}見本市1-1", "06-1234-5678", "開設 太郎", "院長 花子", "令2.1.1\
新規\
令2.1.1", "歯", "診療所\
現存"])
    out = BytesIO(); book.save(out); return out.getvalue()


def test_v13_dental_workbook_parses_as_dental_with_collision_safe_id():
    source = OfficialSource(
        "dental_test", "近畿厚生局", "大阪府・歯科",
        "https://kouseikyoku.mhlw.go.jp/x.xlsx", "https://kouseikyoku.mhlw.go.jp/", "xlsx",
        prefecture_hint="大阪府", medical_type_hint="歯科", expected_prefectures=("大阪府",),
    )
    frame = parse_code_workbook(_dental_record_book(), "2026.9_kikanzentai_osaka_sika.xlsx", source)
    assert len(frame) == 1
    assert frame.iloc[0]["medical_type"] == "歯科"
    assert frame.iloc[0]["clinic_id"] == "大阪府:歯科:0112345"
    assert frame.iloc[0]["facility_type"] == "診療所"


def test_v13_mixed_zip_selects_only_requested_medical_type():
    from src.master.national_maps_queue import parse_official_payload
    payload = _zip_of_named_books({
        "fukuoka_ika.xlsx": _one_record_book("福岡県", "01,2222,2"),
        "fukuoka_shika.xlsx": _dental_record_book("福岡県", "01-33333", "福岡見本歯科"),
    })
    medical = OfficialSource(
        "kyushu_mix_med", "九州厚生局", "福岡県・医科", "https://kouseikyoku.mhlw.go.jp/x.zip",
        "https://kouseikyoku.mhlw.go.jp/", "zip", prefecture_hint="福岡県", medical_type_hint="医科", mixed_archive=True, expected_prefectures=("福岡県",),
    )
    dental = OfficialSource(
        "kyushu_mix_dental", "九州厚生局", "福岡県・歯科", "https://kouseikyoku.mhlw.go.jp/x.zip",
        "https://kouseikyoku.mhlw.go.jp/", "zip", prefecture_hint="福岡県", medical_type_hint="歯科", mixed_archive=True, expected_prefectures=("福岡県",),
    )
    mf = parse_official_payload(payload, medical)
    df = parse_official_payload(payload, dental)
    assert mf["clinic_name"].tolist() == ["福岡県見本クリニック"]
    assert set(mf["medical_type"]) == {"医科"}
    assert df["clinic_name"].tolist() == ["福岡見本歯科"]
    assert set(df["medical_type"]) == {"歯科"}


def test_v13_population_can_filter_dental_only_or_both():
    rows = [
        {"clinic_id":"東京都:0111111", "clinic_name":"見本内科", "address":"東京都千代田区", "phone":"03-1111-1111", "prefecture":"東京都", "medical_type":"医科", "facility_type":"診療所", "status":"現存"},
        {"clinic_id":"東京都:歯科:0111111", "clinic_name":"見本歯科", "address":"東京都千代田区", "phone":"03-2222-2222", "prefecture":"東京都", "medical_type":"歯科", "facility_type":"診療所", "status":"現存"},
    ]
    frame = pd.DataFrame(rows)
    dental, dcounts = build_maps_population(frame, set(), medical_types=("歯科",))
    both, bcounts = build_maps_population(frame, set(), medical_types=("医科","歯科"))
    assert [r["medical_type"] for r in dental] == ["歯科"]
    assert dcounts["non_medical"] == 1
    assert {r["medical_type"] for r in both} == {"医科", "歯科"}
    assert bcounts["target"] == 2


def test_v13_package_reports_medical_and_dental_and_keeps_collector_8_columns():
    rows = []
    for i, medical_type in enumerate(("医科", "歯科"), 1):
        rows.append({
            "clinic_id": f"神奈川県:{'歯科:' if medical_type == '歯科' else ''}01{i:05d}",
            "clinic_name": f"見本{medical_type}{i}", "address": "神奈川県横浜市", "phone": f"045-000-000{i}",
            "prefecture":"神奈川県", "medical_type":medical_type, "facility_type":"診療所", "owner_name":"", "manager_name":"",
            "designation_date":"", "registration_reason":"", "status":"現存", "as_of":"2026-09-01", "source_url":"",
            "designation_period_start":"", "departments":"", "source_bureau":"テスト", "source_file":"x.xlsx",
        })
    frame = pd.DataFrame(rows, columns=MASTER_HEADERS)
    package, manifest, files = build_package(frame, set(), [], {"completed":0}, medical_types=("医科","歯科"))
    assert manifest["medical_types"] == ["医科", "歯科"]
    assert manifest["per_medical_type_target"] == {"医科":1, "歯科":1}
    with zipfile.ZipFile(BytesIO(package)) as zf:
        header = zf.read("maps_queue_all.csv").decode("utf-8-sig").splitlines()[0].split(",")
        assert header == QUEUE_HEADERS
        body = zf.read("maps_queue_all.csv").decode("utf-8-sig")
        assert "神奈川県:歯科:" in body



def _tohoku_collocation_book(primary_type: str, collocation: str, clinic_name: str) -> bytes:
    book = Workbook(); ws = book.active
    ws.title = "青森県"
    ws.append(["コード内容別医療機関一覧表"])
    ws.append([f"[令和 8年 9月 1日現在　{primary_type}　現存/休止]", collocation])
    code = "01-1024-3"
    ws.append(["1", code, clinic_name, "〒030-0000 青森県青森市見本1-1", "017-123-4567", "A", "B", "令2.1.1", "内" if primary_type == "医科" else "歯", "診療所\n現存"])
    out = BytesIO(); book.save(out); return out.getvalue()


def test_v131_dedicated_tohoku_medical_not_rejected_by_dental_collocation_marker():
    source = OfficialSource(
        "tohoku_live_shape", "東北厚生局", "東北6県・医科",
        "https://kouseikyoku.mhlw.go.jp/tohoku/shitei-touhoku-ika-r0809.xlsx",
        "https://kouseikyoku.mhlw.go.jp/tohoku/gyomu/gyomu/hoken_kikan/itiran.html",
        "xlsx", medical_type_hint="医科", expected_prefectures=("青森県",),
    )
    raw = _tohoku_collocation_book("医科", "歯科併設", "青森見本内科")
    from src.master.national_maps_queue import parse_official_payload
    frame = parse_official_payload(raw, source)
    assert len(frame) == 1
    assert frame.iloc[0]["medical_type"] == "医科"
    assert frame.iloc[0]["clinic_name"] == "青森見本内科"


def test_v131_dedicated_tohoku_dental_not_rejected_by_medical_collocation_marker():
    source = OfficialSource(
        "tohoku_dental_live_shape", "東北厚生局", "東北6県・歯科",
        "https://kouseikyoku.mhlw.go.jp/tohoku/shitei-touhoku-shika-r0809.xlsx",
        "https://kouseikyoku.mhlw.go.jp/tohoku/gyomu/gyomu/hoken_kikan/itiran.html",
        "xlsx", medical_type_hint="歯科", expected_prefectures=("青森県",),
    )
    raw = _tohoku_collocation_book("歯科", "医科併設", "青森見本歯科")
    from src.master.national_maps_queue import parse_official_payload
    frame = parse_official_payload(raw, source)
    assert len(frame) == 1
    assert frame.iloc[0]["medical_type"] == "歯科"
    assert frame.iloc[0]["clinic_name"] == "青森見本歯科"


def test_v131_strong_type_detection_ignores_collocation_words():
    from src.master.national_maps_queue import _strong_medical_type_from_text
    assert _strong_medical_type_from_text("令和8年9月1日現在 医科 現存/休止 歯科併設") == "医科"
    assert _strong_medical_type_from_text("令和8年9月1日現在 歯科 現存/休止 医科併設") == "歯科"
    assert _strong_medical_type_from_text("令和8年9月1日現在 薬局 現存/休止") == "薬局"


def test_v131_only_kyushu_registry_sources_are_marked_mixed_archives():
    mixed = [s for s in OFFICIAL_SOURCES if s.mixed_archive]
    assert len(mixed) == 16
    assert all(s.bureau == "九州厚生局" for s in mixed)
    urls = {}
    for source in mixed:
        urls.setdefault(source.url, set()).add(source.medical_type_hint)
    assert len(urls) == 8
    assert all(types == {"医科", "歯科"} for types in urls.values())


def test_v131_cached_parse_failure_redownloads_once(monkeypatch, tmp_path):
    import src.master.national_maps_queue as nmq
    source = OfficialSource(
        "retry_source", "テスト局", "再取得テスト",
        "https://kouseikyoku.mhlw.go.jp/retry.xlsx",
        "https://kouseikyoku.mhlw.go.jp/", "xlsx",
        prefecture_hint="青森県", medical_type_hint="医科", expected_prefectures=("青森県",),
    )
    good = _tohoku_collocation_book("医科", "歯科併設", "再取得見本クリニック")
    calls = []
    def fake_download(src, cache_dir, timeout=90, force=False):
        calls.append(force)
        p = tmp_path / "retry_source.xlsx"
        if force:
            p.write_bytes(good)
            return good, p, False
        bad = b"PK" + b"broken" * 300
        p.write_bytes(bad)
        return bad, p, True
    monkeypatch.setattr(nmq, "download_source", fake_download)
    monkeypatch.setattr(nmq, "PREFECTURES", ("青森県",))
    frame, report = nmq.collect_all_official(tmp_path, sources=(source,))
    assert len(frame) == 1
    assert calls == [False, True]
    assert report[0]["mode"] == "redownload"


def test_v131_registry_current_official_urls_for_dedicated_dental_sources():
    # 2026-09-18時点の各厚生局公式一覧ページで確認した専用歯科ファイル。
    expected = {
        "hokkaido_dental": "https://kouseikyoku.mhlw.go.jp/hokkaido/000499341.xlsx",
        "tohoku_dental": "https://kouseikyoku.mhlw.go.jp/tohoku/shitei-touhoku-shika-r0809.xlsx",
        "kanto_shinetsu_dental": "https://kouseikyoku.mhlw.go.jp/kantoshinetsu/shitei_shika_r0809.zip",
        "tokai_hokuriku_dental": "https://kouseikyoku.mhlw.go.jp/tokaihokuriku/2609-01-03.zip",
        "kinki_dental": "https://kouseikyoku.mhlw.go.jp/kinki/2026.9_kikanzentai_sika.zip",
        "chugoku_dental": "https://kouseikyoku.mhlw.go.jp/chugokushikoku/000500019.zip",
        "shikoku_dental": "https://kouseikyoku.mhlw.go.jp/shikoku/000499485.zip",
    }
    registry = {s.source_id: s.url for s in OFFICIAL_SOURCES}
    for source_id, url in expected.items():
        assert registry[source_id] == url
